import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import json
import traceback
import zipfile
from processor import TemplateProcessor

def reproduce():
    # Load case session data
    session_path = "cases/case_1781428019/session.json"
    if not os.path.exists(session_path):
        print("Session not found")
        return
        
    with open(session_path, "r", encoding="utf-8") as f:
        session = json.load(f)
        
    data = session["data"]
    verified_fields = set(session.get("verified_fields", []))
    
    # Replicate route logic for ds list mapping and pad arrays
    import re
    from collections import defaultdict
    
    sec_sched_val = data.get("second_schedule", "")
    sec_sched_val = re.sub(r'(?<=\S)\s+(Original\b|Certified Copy\b)', r'\n\1', sec_sched_val, flags=re.IGNORECASE)
    sec_lines = [x.strip() for x in sec_sched_val.split("\n") if x.strip()]
    ds_list = [{"t": line} for line in sec_lines]
    while len(ds_list) < 10:
        ds_list.append({"t": ""})
    data["ds"] = ds_list

    for key in ["bs", "ls", "ps", "ws", "ss", "sellers", "buyers", "chain"]:
        if key not in data or not isinstance(data[key], list):
            data[key] = []
        while len(data[key]) < 10:
            data[key].append(defaultdict(str))

    context = data.copy()
    context["w1"] = data["ws"][0]
    context["w2"] = data["ws"][1]
    
    # Generate chain_text dynamically for context
    if "title_chain" in context and isinstance(context["title_chain"], list):
        # In route, _generate_chain_narrative is local, let's copy its logic or use app._generate_chain_narrative
        from app import _generate_chain_narrative
        context["chain_text"] = _generate_chain_narrative(context["title_chain"])
    else:
        context["chain_text"] = ""
        
    context['d'] = context.copy()

    # Try all 4 templates in SALE_DEED
    templates = [
        "Balkishan Yadav & Vinod Gurjar MM PNo 72 Shiv Enclave Vill Gawar Brahmani Sanganer.docx",
        "Ganesh Pareek & hansa Devi MW P NO 35-B Krishna Vihar village BadhShyaopura vatika Road BT.docx",
        "SD-Vivek Saxena,  Sunita Saxena - Vijay Laxmi - JDA+2SD+Flat_unicode (1)_devlys.docx",
        "manoj kumar & Monu kumari MW +EAST PART OF BALAJI NAGAR+SO PATTA+3REC+SD.docx"
    ]
    
    for template in templates:
        template_path = os.path.join("templates/SALE_DEED", template)
        output_path = f"scratch/test_{template}"
        print(f"\n--- Testing template: {template} ---")
        try:
            processor = TemplateProcessor(template_path)
            processor.generate(context, output_path, highlight_ai=True, highlight_missing=True, verified_fields=verified_fields)
            print(f"SUCCESS: Generated {output_path}")
            
            # Check XML validity of the generated zip file (docx is a zip file)
            with zipfile.ZipFile(output_path) as z:
                # Read document.xml
                doc_xml = z.read("word/document.xml")
                # Try parsing it to check if it's well-formed XML
                from xml.etree import ElementTree as ET
                try:
                    ET.fromstring(doc_xml)
                    print("XML parsed successfully! No malformed XML in word/document.xml.")
                    
                    # Try loading it with docx
                    import docx
                    d = docx.Document(output_path)
                    # Iterate through paragraphs and tables to verify no internal errors
                    for p in d.paragraphs:
                        _ = p.text
                        for r in p.runs:
                            _ = r.text
                    for t in d.tables:
                        for row in t.rows:
                            for cell in row.cells:
                                for p in cell.paragraphs:
                                    _ = p.text
                    print("python-docx loaded and read the document successfully!")
                except ET.ParseError as pe:
                    print(f"MALFORMED XML ERROR in word/document.xml: {pe}")
                    # Print context around parse error if possible
                    # Find line/column of error
                    print(traceback.format_exc())
                    
        except Exception as e:
            print(f"FAILED: {e}")
            traceback.print_exc()

if __name__ == "__main__":
    reproduce()
