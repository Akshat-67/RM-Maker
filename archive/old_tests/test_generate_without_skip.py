import sys
import os
import json
import re
from docx import Document

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from template_tools.builder_core import DocManipulator

def count_placeholders(doc):
    count = 0
    tags = []
    for part_name, part in DocManipulator.iter_story_parts(doc):
        for p in DocManipulator.iter_paragraphs_deep(part):
            matches = re.findall(r"\{\{.*?\}\}", p.text)
            if matches:
                count += len(matches)
                tags.extend(matches)
    return count, tags

def main():
    mappings_path = "scratch/fresh_discovery_mappings.json"
    doc_path = "templates/SALE_DEED/SD-Vivek Saxena,  Sunita Saxena - Vijay Laxmi - JDA+2SD+Flat.docx"
    output_path = "scratch/fresh_generated_template.docx"
    
    if not os.path.exists(mappings_path):
        print(f"Error: {mappings_path} not found")
        return
        
    with open(mappings_path, "r", encoding="utf-8") as f:
        mapping_dict = json.load(f)
        
    doc = Document(doc_path)
    
    DocManipulator.apply_deterministic_rules(doc)
    
    reps = sorted(mapping_dict.items(), key=lambda x: len(x[0]), reverse=True)
    for part_name, part in DocManipulator.iter_story_parts(doc):
        for p in DocManipulator.iter_paragraphs_deep(part):
            for old, new in reps:
                DocManipulator.safe_replace(p, old, new)
                
    doc.save(output_path)
    print(f"Generated fresh template at {output_path}")
    
    placeholder_count, tags = count_placeholders(doc)
    print(f"Total placeholders found in generated template: {placeholder_count}")
    print("Unique placeholders:", sorted(list(set(tags))))

if __name__ == "__main__":
    main()
