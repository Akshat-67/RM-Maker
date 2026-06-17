from docx import Document

doc_path = r"C:\Users\aksha\Documents\RM Generator\RM-Maker\templates\SALE_DEED\SD-Vivek Saxena,  Sunita Saxena - Vijay Laxmi - JDA+2SD+Flat_unicode (1)_devlys.docx"
doc = Document(doc_path)

for i, para in enumerate(doc.paragraphs):
    if "chain" in para.text:
        print(f"Found 'chain' in paragraph {i}: {para.text}")

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for i, para in enumerate(cell.paragraphs):
                if "chain" in para.text:
                    print(f"Found 'chain' in table cell paragraph: {para.text}")

