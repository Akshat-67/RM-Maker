import os
import zipfile
import re
from docx import Document
from docx.shared import Pt
from docxtpl import DocxTemplate

TEMPLATE_PATH = "devlys_poc_template.docx"
OUTPUT_PATH = "devlys_poc_output.docx"

def create_poc_template():
    print(f"--- Creating POC Template: {TEMPLATE_PATH} ---")
    doc = Document()
    
    # Add a paragraph with DevLys 040 font
    p = doc.add_paragraph()
    run = p.add_run("This is DevLys 040 text: ")
    run.font.name = 'DevLys 040'
    run.font.size = Pt(12)
    
    # Insert placeholders inside the DevLys run
    run_ph = p.add_run("{{ name }} {{ seller }} {{ amount }}")
    run_ph.font.name = 'DevLys 040'
    run_ph.font.size = Pt(12)
    
    doc.save(TEMPLATE_PATH)
    print("Template created.")

def inspect_xml(path):
    print(f"\n--- Inspecting XML for: {path} ---")
    with zipfile.ZipFile(path) as zf:
        xml_content = zf.read("word/document.xml").decode("utf-8")
        # Find the placeholders in XML
        placeholders = re.findall(r'\{\{.*?\}\}', xml_content)
        print(f"Placeholders found in XML: {placeholders}")
        
        # Check for font definitions
        has_devlys = 'w:rFonts w:ascii="DevLys 040"' in xml_content or 'w:rFonts w:hAnsi="DevLys 040"' in xml_content
        print(f"DevLys 040 font reference found: {has_devlys}")
        
        # Print a snippet of the XML around a placeholder
        ph_match = re.search(r'(<w:r>.*?\{\{.*?\}\}.*?</w:r>)', xml_content)
        if ph_match:
            print(f"XML Snippet: {ph_match.group(1)}")

def test_generation():
    print(f"\n--- Testing docxtpl Generation: {OUTPUT_PATH} ---")
    doc = DocxTemplate(TEMPLATE_PATH)
    context = {
        "name": "Praveen Singh",
        "seller": "Vikram Singh",
        "amount": "15,00,000"
    }
    doc.render(context)
    doc.save(OUTPUT_PATH)
    print("Document generated.")

if __name__ == "__main__":
    create_poc_template()
    inspect_xml(TEMPLATE_PATH)
    test_generation()
    inspect_xml(OUTPUT_PATH)
