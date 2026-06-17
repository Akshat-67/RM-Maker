import os
import sys

# Ensure the root project folder is in the python path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from utils.devlys_to_unicode import DevLysToUnicodeConverter
from docx import Document

def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass
    input_deed = r"templates/SALE_DEED/manoj kumar & Monu kumari MW +EAST PART OF BALAJI NAGAR+SO PATTA+3REC+SD.docx"
    output_deed = r"scratch/converted_manoj_kumar.docx"

    print("Converting DevLys DOCX to Unicode...")
    try:
        DevLysToUnicodeConverter.convert_docx(input_deed, output_deed)
        print(f"Conversion complete! Saved to {output_deed}")
    except Exception as e:
        print(f"Error during conversion: {e}")
        import traceback
        traceback.print_exc()
        return

    # Verify contents
    original_doc = Document(input_deed)
    converted_doc = Document(output_deed)

    print("\n--- Verification of conversion ---")
    print(f"Original paragraphs count: {len(original_doc.paragraphs)}")
    print(f"Converted paragraphs count: {len(converted_doc.paragraphs)}")

    print("\n--- Sample Paragraphs Comparison (First 5 non-empty paragraphs) ---")
    count = 0
    for i, (orig_p, conv_p) in enumerate(zip(original_doc.paragraphs, converted_doc.paragraphs)):
        orig_text = orig_p.text.strip()
        conv_text = conv_p.text.strip()
        if orig_text:
            count += 1
            print(f"\n[Paragraph {i}]")
            print(f"Original : {orig_text}")
            print(f"Converted: {conv_text}")
            if count >= 5:
                break

    # Search for known phrases like 'विक्रय-पत्र' or 'विक्रेता' or 'क्रेता' in converted
    print("\n--- Checking for Devanagari Unicode Keywords ---")
    keywords = ["विक्रय-पत्र", "विक्रय", "क्रेता", "विक्रेता", "विलेख", "श्री"]
    found_keywords = {kw: 0 for kw in keywords}
    for p in converted_doc.paragraphs:
        for kw in keywords:
            if kw in p.text:
                found_keywords[kw] += 1

    for kw, occurrences in found_keywords.items():
        print(f"Keyword '{kw}': found {occurrences} times.")

if __name__ == "__main__":
    main()
