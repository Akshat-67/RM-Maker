import sys
import os
from docx import Document

sys.path.append(r"C:\Users\aksha\Documents\RM Generator\RM-Maker")
from utils.devlys_to_unicode import DevLysToUnicodeConverter

firm_path = r"C:\Users\aksha\Documents\RM Generator\RM-Maker\Reference Materials\Sale Deed\FLAT\SD-Vivek Saxena,  Sunita Saxena - Vijay Laxmi - JDA+2SD+Flat.docx"
gen_path = r"C:\Users\aksha\Documents\RM Generator\RM-Maker\cases\case_1781374115\Generated_AI_Test3.docx"

def convert_if_devlys(text):
    return DevLysToUnicodeConverter.devlys_to_unicode_text(text)

def save_range(filepath, out_path, start, end):
    doc = Document(filepath)
    lines = []
    lines.append(f"=== File: {os.path.basename(filepath)} (Range {start} to {end}) ===")
    
    for i in range(min(start, len(doc.paragraphs)), min(end, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        text = p.text.strip()
        if text:
            converted = convert_if_devlys(text)
            lines.append(f"P[{i}]: {converted}")
            
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(lines))

save_range(firm_path, r"C:\Users\aksha\Documents\RM Generator\RM-Maker\scratch\firm_exact_chain.txt", 0, 26)
save_range(gen_path, r"C:\Users\aksha\Documents\RM Generator\RM-Maker\scratch\generated_exact_chain.txt", 0, 26)

print("Saved exact chain ranges!")
