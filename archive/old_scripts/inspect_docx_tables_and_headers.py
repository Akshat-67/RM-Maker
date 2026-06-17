import sys
import os
from docx import Document

def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass

    input_deed = r"templates/SALE_DEED/manoj kumar & Monu kumari MW +EAST PART OF BALAJI NAGAR+SO PATTA+3REC+SD.docx"
    output_deed = r"scratch/converted_manoj_kumar.docx"

    if not os.path.exists(output_deed):
        print("Converted deed not found.")
        return

    orig_doc = Document(input_deed)
    conv_doc = Document(output_deed)

    print("=== TABLES VERIFICATION ===")
    print(f"Original tables count: {len(orig_doc.tables)}")
    print(f"Converted tables count: {len(conv_doc.tables)}")

    for t_idx, (orig_t, conv_t) in enumerate(zip(orig_doc.tables, conv_doc.tables)):
        print(f"\nTable {t_idx} (rows={len(orig_t.rows)}, cols={len(orig_t.columns)})")
        # Print first cell text from first row
        if orig_t.rows:
            orig_cell = orig_t.rows[0].cells[0].text.strip()
            conv_cell = conv_t.rows[0].cells[0].text.strip()
            print(f"  Row 0 Cell 0 Original : {repr(orig_cell)}")
            print(f"  Row 0 Cell 0 Converted: {repr(conv_cell)}")

    print("\n=== HEADERS & FOOTERS VERIFICATION ===")
    for s_idx, (orig_sec, conv_sec) in enumerate(zip(orig_doc.sections, conv_doc.sections)):
        print(f"\nSection {s_idx}:")
        for hf_name in ["header", "footer"]:
            orig_hf = getattr(orig_sec, hf_name)
            conv_hf = getattr(conv_sec, hf_name)
            
            orig_text = "\n".join(p.text for p in orig_hf.paragraphs).strip() if orig_hf else ""
            conv_text = "\n".join(p.text for p in conv_hf.paragraphs).strip() if conv_hf else ""
            
            if orig_text:
                print(f"  {hf_name.upper()} Original : {repr(orig_text)}")
                print(f"  {hf_name.upper()} Converted: {repr(conv_text)}")

if __name__ == "__main__":
    main()
