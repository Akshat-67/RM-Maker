from docx import Document

doc_path = r"C:\Users\aksha\Documents\RM Generator\RM-Maker\templates\SALE_DEED\SD-Vivek Saxena,  Sunita Saxena - Vijay Laxmi - JDA+2SD+Flat_unicode (1)_devlys.docx"
doc = Document(doc_path)

found = False
for i, para in enumerate(doc.paragraphs):
    if "chain_text" in para.text:
        print(f"Found in paragraph {i}: {para.text}")
        found = True

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for i, para in enumerate(cell.paragraphs):
                if "chain_text" in para.text:
                    print(f"Found in table cell paragraph: {para.text}")
                    found = True

if not found:
    print("Not found in paragraphs or tables!")
