"""Dump the raw XML of converted runs, document defaults, and styles to find the font cascade issue."""
import sys, os
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
try: sys.stdout.reconfigure(encoding='utf-8')
except: pass

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

def pp(el):
    return etree.tostring(el, pretty_print=True).decode('utf-8')

def main():
    doc = Document("scratch/converted_manoj_kumar.docx")

    # 1. Dump docDefaults
    styles_part = doc.part.element.find(qn('w:body'))
    print("=== DOCUMENT DEFAULTS (from styles.xml) ===")
    try:
        styles_element = doc.styles.element
        doc_defaults = styles_element.find(qn('w:docDefaults'))
        if doc_defaults is not None:
            print(pp(doc_defaults))
        else:
            print("No docDefaults found")
    except Exception as e:
        print(f"Error: {e}")

    # 2. Dump the "Normal" style
    print("\n=== NORMAL STYLE ===")
    try:
        for style_el in doc.styles.element.findall(qn('w:style')):
            style_id = style_el.get(qn('w:styleId'))
            if style_id == 'Normal':
                print(pp(style_el))
                break
    except Exception as e:
        print(f"Error: {e}")

    # 3. Dump first converted paragraph's style + first run XML
    print("\n=== PARAGRAPH 22 STYLE & RUN XML ===")
    p = doc.paragraphs[22]  # "विक्रय-पत्र"
    print(f"Text: {p.text}")
    print(f"Style: {p.style.name if p.style else 'None'}")

    # paragraph XML properties
    pPr = p._p.find(qn('w:pPr'))
    if pPr is not None:
        print(f"Paragraph Properties:\n{pp(pPr)}")

    # first run's full XML
    if p.runs:
        r = p.runs[0]
        print(f"\nRun 0 full XML:\n{pp(r._r)}")

    # 4. Check the style that paragraph 22 uses
    print("\n=== PARAGRAPH 22's STYLE DEFINITION ===")
    if p.style:
        print(pp(p.style.element))

    # 5. Also check original document's defaults for comparison
    print("\n=== ORIGINAL DOC DEFAULTS ===")
    orig = Document("templates/SALE_DEED/manoj kumar & Monu kumari MW +EAST PART OF BALAJI NAGAR+SO PATTA+3REC+SD.docx")
    try:
        orig_defaults = orig.styles.element.find(qn('w:docDefaults'))
        if orig_defaults is not None:
            print(pp(orig_defaults))
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    main()
