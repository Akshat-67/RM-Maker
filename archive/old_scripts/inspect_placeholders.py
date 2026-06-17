"""Quickly inspect Jinja2 placeholders in a word template to see how r/rn are used."""
import sys, os
from docx import Document
from lxml import etree
import zipfile
import re

def main():
    template_path = "templates/SALE_DEED/manoj kumar & Monu kumari MW +EAST PART OF BALAJI NAGAR+SO PATTA+3REC+SD.docx"
    if not os.path.exists(template_path):
        print(f"Error: Template {template_path} not found")
        return

    doc = Document(template_path)
    print("=== FIRST 40 PARAGRAPHS ===")
    for idx, p in enumerate(doc.paragraphs[:40]):
        print(f"P{idx:3d}: {p.text}")

if __name__ == "__main__":
    main()
