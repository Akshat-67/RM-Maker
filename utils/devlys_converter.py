import os
import re
import zipfile
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from html import unescape

def unicode_to_devlys(text):
    if not text:
        return ""

    array_one = [
        "‘", "’", "“", "”", "(", ")", "{", "}", "=", "।", "?", "-", "µ", ",", ".", "् ",
        "०", "१", "२", "३", "४", "५", "६", "७", "८", "९", "x",
        "फ़्", "क़", "ख़", "ग़", "ज़्", "ज़", "ड़", "ढ़", "फ़", "य़", "ऱ", "ऩ",
        "त्त्", "त्त", "क्त", "दृ", "कृ",
        "ह्न", "ह्य", "हृ", "ह्म", "ह्र", "ह्", "द्द", "क्ष्", "क्ष", "त्र्", "त्र", "ज्ञ",
        "छ्य", "ट्य", "ठ्य", "ड्य", "ढ्य", "द्य", "द्व",
        "श्र", "ट्र", "ड्र", "ढ्र", "छ्र", "क्र", "फ्र", "द्र", "प्र", "ग्र", "रु", "रू",
        "्र",
        "ओ", "औ", "आ", "अ", "ई", "इ", "उ", "ऊ", "ऐ", "ए", "ऋ",
        "क्", "क", "क्क", "ख्", "ख", "ग्", "ग", "घ्", "घ", "ङ",
        "चै", "च्", "च", "छ", "ज्", "ज", "झ्", "झ", "ञ",
        "ट्ट", "ट्ठ", "ट", "ठ", "ड्ड", "ड्ढ", "ड", "ढ", "ण्", "ण",
        "त्", "त", "थ्", "थ", "द्ध", "द", "ध्", "ध", "न्", "न",
        "प्", "प", "फ्", "फ", "ब्", "ब", "भ्", "भ", "म्", "म",
        "य्", "य", "र", "ल्", "ल", "ळ", "व्", "व",
        "श्", "श", "ष्", "ष", "स्", "स", "ह",
        "ऑ", "ॉ", "ो", "ौ", "ा", "ी", "ु", "ू", "ृ", "े", "ै",
        "ं", "ँ", "ः", "ॅ", "ऽ", "् ", "्"
    ]

    array_two = [
        "^", "*", "Þ", "ß", "¼", "½", "¿", "À", "¾", "A", "\\", "&", "&", "]", "-", " ",
        "å", "ƒ", "„", "…", "†", "‡", "ˆ", "‰", "Š", "‹", "Û",
        "¶", "d", "[k", "x", "T", "t", "M+", "<+", "Q", ";", "j", "u",
        "Ù", "Ùk", "ä", "–", "—",
        "à", "á", "â", "ã", "ºz", "º", "í", "{", "{k", "«", "=", "K",
        "Nî", "Vî", "Bî", "Mî", "<î", "|", "}",
        "J", "Vª", "Mª", "<ªª", "Nª", "Ø", "Ý", "æ", "ç", "xz", "#", ":",
        "z",
        "vks", "vkS", "vk", "v", "bZ", "b", "m", "Å", ",s", ",", "_",
        "D", "d", "ô", "[", "[k", "X", "x", "?", "?k", "³",
        "pkS", "P", "p", "N", "T", "t", "÷", ">", "¥",
        "ê", "ë", "V", "B", "ì", "ï", "M", "<", ".", ".k",
        "R", "r", "F", "Fk", ")", "n", "/", "/k", "U", "u",
        "I", "i", "¶", "Q", "C", "c", "H", "Hk", "E", "e",
        "¸", ";", "j", "Y", "y", "G", "O", "o",
        "'", "'k", "\"", "\"k", "L", "l", "g",
        "v‚", "‚", "ks", "kS", "k", "h", "q", "w", "`", "s", "S",
        "a", "¡", "%", "W", "·", " ", "~"
    ]

    modified_substring = text

    # Remove the Unicode abbreviation dot (U+0970) as it causes rendering artifacts in DevLys
    modified_substring = modified_substring.replace("॰", "")

    # Initial normalization for nuqta characters
    modified_substring = modified_substring.replace("क़", "क़")
    modified_substring = modified_substring.replace("ख़", "ख़")
    modified_substring = modified_substring.replace("ग़", "ग़")
    modified_substring = modified_substring.replace("ज़", "ज़")
    modified_substring = modified_substring.replace("ड़", "ड़")
    modified_substring = modified_substring.replace("ढ़", "ढ़")
    modified_substring = modified_substring.replace("ऩ", "ऩ")
    modified_substring = modified_substring.replace("फ़", "फ़")
    modified_substring = modified_substring.replace("य़", "य़")
    modified_substring = modified_substring.replace("ऱ", "ऱ")

    # Handle the "i" matra (ि)
    # It needs to move to the front of the character/conjunct
    position_of_f = modified_substring.find("ि")
    while position_of_f != -1:
        if position_of_f > 0:
            character_left_to_f = modified_substring[position_of_f - 1]
            modified_substring = modified_substring[:position_of_f-1] + "f" + character_left_to_f + modified_substring[position_of_f+1:]
            
            curr_pos = position_of_f - 1
            while curr_pos > 0 and modified_substring[curr_pos-1] == "्":
                # It's a conjunct, move "f" further left
                string_to_be_replaced = modified_substring[curr_pos-2] + "्"
                modified_substring = modified_substring[:curr_pos-2] + "f" + string_to_be_replaced + modified_substring[curr_pos+1:]
                curr_pos = curr_pos - 2
            
            position_of_f = modified_substring.find("ि", position_of_f + 1)
        else:
            position_of_f = modified_substring.find("ि", position_of_f + 1)

    # Handle half-R (र्)
    set_of_matras = "ािीुूृेैोौं:ँॅ"
    modified_substring += "  " # Buffer as per JS

    position_of_half_R = modified_substring.find("र्")
    while position_of_half_R != -1:
        prob_pos = position_of_half_R + 2
        
        while prob_pos < len(modified_substring):
            if prob_pos + 1 < len(modified_substring) and modified_substring[prob_pos+1] in set_of_matras:
                prob_pos += 1
            else:
                break
        
        target = modified_substring[position_of_half_R+2 : prob_pos+1]
        modified_substring = modified_substring[:position_of_half_R] + target + "Z" + modified_substring[prob_pos+1:]
        position_of_half_R = modified_substring.find("र्")

    modified_substring = modified_substring[:-2] # Remove buffer

    # Simple character replacements
    for i in range(len(array_one)):
        modified_substring = modified_substring.replace(array_one[i], array_two[i])

    return modified_substring

class UnicodeToDevLysConverter:
    @staticmethod
    def is_hindi(text):
        return any("\u0900" <= char <= "\u097F" for char in text)

    @staticmethod
    def set_font_devlys(run, font_name="DevLys 040"):
        run.font.name = font_name
        try:
            rPr = run._r.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:ascii'), font_name)
            rFonts.set(qn('w:hAnsi'), font_name)
            # Do NOT set w:cs or w:hint="cs" for DevLys, as it's an ASCII font
        except Exception:
            pass

    @staticmethod
    def set_font_arial(run, font_name="Arial"):
        run.font.name = font_name
        try:
            rPr = run._r.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:ascii'), font_name)
            rFonts.set(qn('w:hAnsi'), font_name)
            # Remove CS hint so Word treats it as standard Latin
            rPr.xpath('./w:rFonts')[0].attrib.pop(qn('w:hint'), None)
        except Exception:
            pass

    @classmethod
    def convert_docx(cls, input_docx_path, output_docx_path, font_name="DevLys 040"):
        if not os.path.exists(input_docx_path):
            raise FileNotFoundError(f"Input file not found: {input_docx_path}")

        doc = Document(input_docx_path)

        def process_paragraph_robust(p):
            # Combine all text first to simplify splitting
            # Note: This approach clears original formatting if it varies within a paragraph.
            # But for template builders, usually the whole paragraph is one style.
            
            full_text = p.text
            if not full_text: return
            if not cls.is_hindi(full_text) and "{{" not in full_text: return

            # Split by Jinja2 placeholders
            parts = re.split(r'(\{\{.*?\}\})', full_text)
            
            # Save original run formatting from first run if it exists
            bold, italic, underline, size, color = None, None, None, None, None
            if p.runs:
                r = p.runs[0]
                bold, italic, underline = r.bold, r.italic, r.underline
                size = r.font.size
                color = r.font.color.rgb

            # Clear existing runs
            for run in p.runs:
                run.text = ""
            
            # Fill with new formatted runs
            for part in parts:
                if not part: continue
                new_run = p.add_run()
                
                # Restore formatting
                new_run.bold = bold
                new_run.italic = italic
                new_run.underline = underline
                if size: new_run.font.size = size
                if color: new_run.font.color.rgb = color

                if re.match(r'\{\{.*?\}\}', part):
                    # Placeholder: Keep Unicode, force Arial
                    new_run.text = part
                    cls.set_font_arial(new_run)
                elif cls.is_hindi(part):
                    # Hindi Text: Convert to DevLys, set DevLys font
                    new_run.text = unicode_to_devlys(part)
                    cls.set_font_devlys(new_run, font_name)
                else:
                    # English/Other: Keep as is
                    new_run.text = part

        for p in doc.paragraphs:
            process_paragraph_robust(p)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        process_paragraph_robust(p)
        for section in doc.sections:
            for h_f in [section.header, section.footer, section.first_page_header, section.first_page_footer, section.even_page_header, section.even_page_footer]:
                if not h_f: continue
                for p in h_f.paragraphs: process_paragraph_robust(p)
                for table in h_f.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs: process_paragraph_robust(p)

        os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)), exist_ok=True)
        doc.save(output_docx_path)
        return output_docx_path

if __name__ == "__main__":
    test_text = "प्रवीन सिंह"
    converted = unicode_to_devlys(test_text)
    print(f"Unicode: {test_text}")
    print(f"DevLys: {converted}")
