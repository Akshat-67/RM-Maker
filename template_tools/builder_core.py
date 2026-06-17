import re
import json
from copy import deepcopy
from docx import Document
from docx.document import Document as _Document
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.text.run import Run

# Core Catalogs
RM_FIELDS = [
    ("RM Execution Date", "{{rd}}"),
    ("Loan Agreement Date", "{{ad}}"),
    ("Borrower 1 - Salutation", "{{bs[0].s}}"),
    ("Borrower 1 - Name", "{{bs[0].n}}"),
    ("Borrower 1 - Age", "{{bs[0].a}}"),
    ("Borrower 1 - Relation", "{{bs[0].r}}"),
    ("Borrower 1 - Rel Name", "{{bs[0].rn}}"),
    ("Borrower 1 - Address", "{{bs[0].adr}}"),
    ("Borrower 1 - Aadhar/ID", "{{bs[0].id}}"),
    ("Borrower 1 - PAN Card", "{{bs[0].pan}}"),
    ("Borrower 2 - Salutation", "{{bs[1].s}}"),
    ("Borrower 2 - Name", "{{bs[1].n}}"),
    ("Borrower 2 - Age", "{{bs[1].a}}"),
    ("Borrower 2 - Relation", "{{bs[1].r}}"),
    ("Borrower 2 - Rel Name", "{{bs[1].rn}}"),
    ("Borrower 2 - Address", "{{bs[1].adr}}"),
    ("Borrower 2 - Aadhar/ID", "{{bs[1].id}}"),
    ("Borrower 2 - PAN Card", "{{bs[1].pan}}"),
    ("Loan 1 - LAN No", "{{ls[0].n}}"),
    ("Loan 1 - Amount (Figures)", "{{ls[0].a}}"),
    ("Loan 1 - Amount (Words)", "{{ls[0].w}}"),
    ("Loan 1 - Tenure", "{{ls[0].t}}"),
    ("Loan 2 - LAN No", "{{ls[1].n}}"),
    ("Loan 2 - Amount (Figures)", "{{ls[1].a}}"),
    ("Loan 2 - Amount (Words)", "{{ls[1].w}}"),
    ("Loan 2 - Tenure", "{{ls[1].t}}"),
    ("Property 1 - Address", "{{ps[0].adr}}"),
    ("Property 1 - North", "{{ps[0].n}}"),
    ("Property 1 - South", "{{ps[0].s}}"),
    ("Property 1 - East", "{{ps[0].e}}"),
    ("Property 1 - West", "{{ps[0].w}}"),
    ("Property 2 - Address", "{{ps[1].adr}}"),
    ("Property 2 - North", "{{ps[1].n}}"),
    ("Property 2 - South", "{{ps[1].s}}"),
    ("Property 2 - East", "{{ps[1].e}}"),
    ("Property 2 - West", "{{ps[1].w}}"),
    ("Bank Signatory - Name", "{{bsign.n}}"),
    ("Bank Signatory - Age", "{{bsign.a}}"),
    ("Bank Signatory - Relation", "{{bsign.r}}"),
    ("Bank Signatory - Rel Name", "{{bsign.rn}}"),
    ("Bank Signatory - PAN Card", "{{bsign.pan}}"),
    ("Bank Signatory - Aadhar/ID", "{{bsign.id}}"),
    ("Witness 1 - Name", "{{ws[0].n}}"),
    ("Witness 1 - Relation", "{{ws[0].r}}"),
    ("Witness 1 - Rel Name", "{{ws[0].rn}}"),
    ("Witness 1 - Address", "{{ws[0].adr}}"),
    ("Witness 2 - Name", "{{ws[1].n}}"),
    ("Witness 2 - Relation", "{{ws[1].r}}"),
    ("Witness 2 - Rel Name", "{{ws[1].rn}}"),
    ("Witness 2 - Address", "{{ws[1].adr}}"),
    ("Document Schedule / Title Chain (FIRST SCHEDULE)", "{{ds_text}}"),
    ("Document Schedule 1 - Title Deed", "{{ds[0].t}}"),
    ("Second Schedule (Documents to be collected)", "{{second_schedule}}"),
]

SD_GENERATED_FIELDS = [
    ("Title Chain Narrative", "{{chain_text}}"),
    ("Property - Full Address", "{{ps[0].full_address}}"),
    ("Property - Dimensions", "{{ps[0].dimension_text}}"),
    ("Property - Boundaries", "{{ps[0].boundary_text}}"),
]

SD_FIELDS = [
    ("Sale Date", "{{rd}}"),
    ("Total Sale Amount", "{{amount}}"),
    ("Sale Amount in Words", "{{amount_words}}"),
    ("Hypothecation (Bank)", "{{hypothecation}}"),
    ("Seller 1 - Name", "{{ss[0].n}}"),
    ("Seller 1 - Age", "{{ss[0].a}}"),
    ("Seller 1 - Caste", "{{ss[0].c}}"),
    ("Seller 1 - Relation", "{{ss[0].r}}"),
    ("Seller 1 - Rel Name", "{{ss[0].rn}}"),
    ("Seller 1 - Address", "{{ss[0].adr}}"),
    ("Seller 1 - Aadhar", "{{ss[0].id}}"),
    ("Seller 1 - PAN", "{{ss[0].pan}}"),
    ("Buyer 1 - Name", "{{bs[0].n}}"),
    ("Buyer 1 - Age", "{{bs[0].a}}"),
    ("Buyer 1 - Caste", "{{bs[0].c}}"),
    ("Buyer 1 - Relation", "{{bs[0].r}}"),
    ("Buyer 1 - Rel Name", "{{bs[0].rn}}"),
    ("Buyer 1 - Address", "{{bs[0].adr}}"),
    ("Buyer 1 - Aadhar", "{{bs[0].id}}"),
    ("Buyer 1 - PAN", "{{bs[0].pan}}"),
    ("Property - Address", "{{ps[0].adr}}"),
    ("Property - Plot No", "{{ps[0].plot_no}}"),
    ("Property - Scheme", "{{ps[0].scheme}}"),
    ("Property - Village", "{{ps[0].village}}"),
    ("Property - Tehsil", "{{ps[0].tehsil}}"),
    ("Property - District", "{{ps[0].dist}}"),
    ("Property - Land Area", "{{ps[0].land_area}}"),
    ("Property - Const Area", "{{ps[0].const_area}}"),
    ("Property - Unit", "{{ps[0].unit}}"),
    ("Property - North", "{{ps[0].n}}"),
    ("Property - South", "{{ps[0].s}}"),
    ("Property - East", "{{ps[0].e}}"),
    ("Property - West", "{{ps[0].w}}"),
    ("Witness 1 - Name", "{{ws[0].n}}"),
    ("Witness 1 - Relation", "{{ws[0].r}}"),
    ("Witness 1 - Rel Name", "{{ws[0].rn}}"),
    ("Witness 1 - Address", "{{ws[0].adr}}"),
    ("Title 1 - Owner", "{{title_chain[0].owner}}"),
    ("Title 1 - Deed Type", "{{title_chain[0].deed_type}}"),
    ("Title 1 - Date", "{{title_chain[0].date}}"),
    ("Title 1 - Book", "{{title_chain[0].book}}"),
    ("Title 1 - Vol", "{{title_chain[0].vol}}"),
    ("Title 1 - Page", "{{title_chain[0].page}}"),
    ("Title 1 - Reg No", "{{title_chain[0].reg_no}}"),
    ("Title 1 - Add Book", "{{title_chain[0].add_book}}"),
    ("Reg - Office", "{{reg.office}}"),
    ("Reg - Book", "{{reg.book}}"),
    ("Reg - Vol", "{{reg.vol}}"),
    ("Reg - Page", "{{reg.page}}"),
    ("Reg - No", "{{reg.reg_no}}"),
    ("Reg - Date", "{{reg.reg_date}}"),
]

class DocManipulator:
    @staticmethod
    def get_doc_content(doc):
        print("[DIAGNOSTIC] DocManipulator.get_doc_content: Starting extraction...")
        chunks = []
        try:
            for part_name, part in DocManipulator.iter_story_parts(doc):
                print(f"[DIAGNOSTIC] DocManipulator.get_doc_content: Processing part {part_name}")
                for p in DocManipulator.iter_paragraphs_deep(part):
                    if p.text.strip():
                        chunks.append(p.text.strip())
        except Exception as e:
            print(f"[DIAGNOSTIC] DocManipulator.get_doc_content: Error during deep iteration: {e}")
            raise
            
        content = "\n".join(chunks)
        print(f"[DIAGNOSTIC] DocManipulator.get_doc_content: Extracted {len(chunks)} chunks, total length {len(content)}")
        return content

    @staticmethod
    def iter_block_items(parent):
        if isinstance(parent, _Document):
            parent_elm = parent.element.body
        elif isinstance(parent, _Cell):
            parent_elm = parent._tc
        else:
            parent_elm = parent._element

        for child in parent_elm.iterchildren():
            if child.tag.endswith('}p'):
                yield Paragraph(child, parent)
            elif child.tag.endswith('}tbl'):
                yield Table(child, parent)

    @staticmethod
    def iter_paragraphs_deep(parent):
        for block in DocManipulator.iter_block_items(parent):
            if isinstance(block, Paragraph):
                yield block
            elif isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        yield from DocManipulator.iter_paragraphs_deep(cell)

    @staticmethod
    def iter_story_parts(doc):
        yield "body", doc
        for section in doc.sections:
            for hf_attr in ['header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer']:
                hf = getattr(section, hf_attr, None)
                if hf:
                    yield hf_attr, hf

    @staticmethod
    def find_fuzzy_match(p, needle):
        if not p.runs: return None
        chars = []
        pos = []
        for ri, run in enumerate(p.runs):
            for ci, c in enumerate(run.text):
                chars.append(" " if c.isspace() or c == "\u00A0" else c)
                pos.append((ri, ci))
        
        haystack = "".join(chars)
        target = re.sub(r"[\s\u00A0]+", " ", needle).strip().casefold()
        pattern = re.escape(target).replace(r"\ ", r"[\s\u00A0]+")
        
        try:
            m = re.search(pattern, haystack, re.IGNORECASE)
            if m:
                s = m.start()
                e = m.end()
                sr, so = pos[s]
                er, eo = pos[e-1]
                return sr, so, er, eo + 1
        except:
            pass

        # Try fast block-matching fuzzy fallback for keys with minor typos
        if len(target) >= 15:
            import difflib
            s_match = difflib.SequenceMatcher(None, target, haystack)
            blocks = s_match.get_matching_blocks()
            best_ratio = 0
            best_span = None
            threshold = 0.85
            
            for block in blocks:
                if block.size < 5:
                    continue
                h_start = max(0, block.b - block.a - 5)
                h_end = min(len(haystack), block.b - block.a + len(target) + 5)
                
                for w_size in range(len(target) - 4, len(target) + 5):
                    for i in range(h_start, min(h_end - w_size + 1, len(haystack))):
                        sub = haystack[i:i+w_size].casefold()
                        s_sub = difflib.SequenceMatcher(None, target, sub)
                        if s_sub.quick_ratio() >= threshold:
                            ratio = s_sub.ratio()
                            if ratio > best_ratio and ratio >= threshold:
                                best_ratio = ratio
                                best_span = (i, i + w_size)
                                
            if best_span:
                s, e = best_span
                sr, so = pos[s]
                er, eo = pos[e-1]
                return sr, so, er, eo + 1
                
        return None

    @staticmethod
    def safe_replace(p, old, new):
        if not old or old == new: return
        match = DocManipulator.find_fuzzy_match(p, old)
        if match:
            s_r, s_o, e_r, e_o = match
            is_tag = "{{" in new
            
            if s_r == e_r:
                run = p.runs[s_r]
                if is_tag:
                    orig_text = run.text
                    prefix = orig_text[:s_o]
                    suffix = orig_text[e_o:]
                    
                    run.text = prefix
                    
                    tag_el = deepcopy(run._element)
                    run._element.addnext(tag_el)
                    tag_run = Run(tag_el, run._parent)
                    tag_run.text = new
                    tag_run.font.name = "Arial"
                    tag_run._element.rPr.get_or_add_rFonts().set(qn('w:ascii'), 'Arial')
                    tag_run._element.rPr.get_or_add_rFonts().set(qn('w:hAnsi'), 'Arial')
                    tag_run._element.rPr.get_or_add_rFonts().set(qn('w:cs'), 'Arial')
                    
                    suffix_el = deepcopy(run._element)
                    tag_el.addnext(suffix_el)
                    suffix_run = Run(suffix_el, run._parent)
                    suffix_run.text = suffix
                else:
                    run.text = run.text[:s_o] + new + run.text[e_o:]
            else:
                p.runs[s_r].text = p.runs[s_r].text[:s_o] + new
                for i in range(s_r + 1, e_r):
                    p.runs[i].text = ""
                p.runs[e_r].text = p.runs[e_r].text[e_o:]
                
                if is_tag:
                    p.runs[s_r].font.name = "Arial"
                    try:
                        p.runs[s_r]._element.rPr.get_or_add_rFonts().set(qn('w:ascii'), 'Arial')
                        p.runs[s_r]._element.rPr.get_or_add_rFonts().set(qn('w:hAnsi'), 'Arial')
                    except: pass

    @staticmethod
    def apply_deterministic_rules(doc):
        """
        Surgically replaces large legal sections with modern narrative tags 
        using high-confidence keyword anchors (supporting both Unicode and DevLys).
        """
        boundary_anchors = [
            "Boundaries as under", "जिसकी चारों सीमाएं", "सीमाऐं निम्न प्रकार", "Bounded as follows",
            "ftldh pkjksa lhek,a", "lhek,sa fuEu çdkj", "lhek,sa fuEu izdkj", "pkjksa lhekvkssa", "pkjksa lhek", "lhekvkssa"
        ]
        dim_anchors = [
            "East to West", "North to South", "Measuring", "जिसकी नाप", "पूर्व से पश्चिम", "उत्तर से दक्षिण",
            "ftldh uki", "iwoZ ls if'pe", "iwoZ ls if’pe", "mRrj ls nf{k.k", "mÙkj ls nf{k.k"
        ]
        addr_anchors = [
            "All that part and parcel", "Situated at", "Municipal No", "Ward No", "स्थित", "प्लॉट नं.", "फ्लैट नं.",
            "fLFkr", "Iy‚V ua-", "IykV ua-", "IykWV ua-", "¶ySV ua-", "¶ysV ua-", "vkoklh; ¶ySV uEcj", "vkoklh; IykV"
        ]
        chain_anchors = [
            "History of Title", "Chain of Ownership", "यह कि उक्त वर्णित संपत्ति का मूल", "History of the property", "Follow of title",
            ";g fd mä of.kZr laifÙk dk ewy", "rRi'pkr~", "rRi’pkr~"
        ]

        # Stop words in Unicode and DevLys
        stop_words = [
            "boundary", " चारों सीमा", "नाप", "पूर्व :", "पश्चिम :", "उत्तर :", "दक्षिण :", "अनुसूची", "schedule", 
            "गवाह", "witness", "हस्ताक्षर", "signature", "शर्त", "term", "क्रेता", "विक्रेता", "flat", "plot", "फ्लैट", "प्लॉट",
            "lhek", "uki", "iwoZ", "if'pe", "if’pe", "mRrj", "nf{k.k", "vuqlwph", "xokg", "gLrk{kj", "Øsrk", "foØsrk"
        ]

        # Person keywords to exclude from property heuristics
        person_keywords = [
            "foØsrk", "fosrk", "foØsrkx.k", "fosrkx.k", "Øsrh", "Øsrk", "Øsrhx.k",
            "xokg", "witness", "tkfr", "fuoklh", "iq= Jh", "iq= Lo-", "iRuh Jh",
            "विक्रेता", "क्रेता", "गवाह", "जाति", "निवासी", "पुत्र", "पत्नी"
        ]

        for part_name, part in DocManipulator.iter_story_parts(doc):
            paragraphs = list(DocManipulator.iter_paragraphs_deep(part))
            
            # Track which tags we've already inserted in this part to avoid duplicates
            inserted_tags = set()
            
            i = 0
            while i < len(paragraphs):
                p = paragraphs[i]
                text = p.text.strip()
                if not text:
                    i += 1
                    continue
                
                # 1. Chain Section Heuristic
                if "{{chain_text}}" not in inserted_tags and any(anchor.lower() in text.lower() for anchor in chain_anchors):
                    DocManipulator.safe_replace(p, text, "{{chain_text}}")
                    inserted_tags.add("{{chain_text}}")
                    
                    # Clear subsequent paragraphs in the chain section to prevent old hardcoded text from leaking.
                    next_idx = i + 1
                    while next_idx < len(paragraphs):
                        next_p = paragraphs[next_idx]
                        next_text = next_p.text.strip()
                        
                        if next_text:
                            if any(anchor.lower() in next_text.lower() for anchor in boundary_anchors + dim_anchors + addr_anchors + stop_words):
                                break
                            
                        # Stop if paragraph is inside a table cell
                        if type(next_p._parent).__name__ == '_Cell':
                            break
                            
                        for run in next_p.runs:
                            run.text = ""
                        next_idx += 1

                # Property narrative heuristics are only run if this paragraph does not describe a person
                is_person = any(k.lower() in text.lower() for k in person_keywords)
                if not is_person:
                    # 2. Boundary Block Heuristic (Surgical first, then Full)
                    if "{{ps[0].boundary_text}}" not in inserted_tags:
                        boundary_match = re.search(r'(ftldh pkjksa lhekvkssa.*?fLFkr gS|जिसकी चारों सीमाएं.*?स्थित है|ftldh pkjksa lhekvkssa.*?dgk x;k gS)', p.text)
                        if boundary_match:
                            DocManipulator.safe_replace(p, boundary_match.group(0), "{{ps[0].boundary_text}}")
                            inserted_tags.add("{{ps[0].boundary_text}}")
                        elif any(anchor.lower() in text.lower() for anchor in boundary_anchors):
                            DocManipulator.safe_replace(p, text, "{{ps[0].boundary_text}}")
                            inserted_tags.add("{{ps[0].boundary_text}}")
                    
                    # 3. Dimension Sentence Heuristic (Surgical first, then Full)
                    if "{{ps[0].dimension_text}}" not in inserted_tags:
                        dim_match = re.search(r'(ftldh uki.*?oxZxt gS|जिसकी नाप.*?वर्गगज है|ftldh uki.*?oxZQhV gS|जिसकी नाप.*?वर्गफीट है)', p.text)
                        if dim_match:
                            DocManipulator.safe_replace(p, dim_match.group(0), "{{ps[0].dimension_text}}")
                            inserted_tags.add("{{ps[0].dimension_text}}")
                        elif any(anchor.lower() in text.lower() for anchor in dim_anchors):
                            DocManipulator.safe_replace(p, text, "{{ps[0].dimension_text}}")
                            inserted_tags.add("{{ps[0].dimension_text}}")

                    # 4. Property Address Heuristic
                    if "{{ps[0].full_address}}" not in inserted_tags and any(anchor.lower() in text.lower() for anchor in addr_anchors):
                        if any(x in text for x in ["Village", "Tehsil", "District", "ग्राम", "तहसील", "जिला", "राजस्थान", "xzke", "t;iqj", "jktLFkku"]):
                            DocManipulator.safe_replace(p, text, "{{ps[0].full_address}}")
                            inserted_tags.add("{{ps[0].full_address}}")

                i += 1

def generate_master_template(doc_path, mapping_dict, output_path):
    doc = Document(doc_path)
    
    # First: Apply Deterministic Heuristics for large sections
    DocManipulator.apply_deterministic_rules(doc)
    
    # Second: Apply AI-discovered entity mappings for short fields
    reps = sorted(mapping_dict.items(), key=lambda x: len(x[0]), reverse=True)
    
    for part_name, part in DocManipulator.iter_story_parts(doc):
        for p in DocManipulator.iter_paragraphs_deep(part):
            for old, new in reps:
                DocManipulator.safe_replace(p, old, new)
    
    doc.save(output_path)

def clean_mapping(mapping):
    if not isinstance(mapping, dict): return {}
    clean = {}
    for k, v in mapping.items():
        if k and v and "{{" in str(v):
            clean[str(k).strip()] = str(v).strip()
    return clean

def get_discovery_prompt(content, mode):
    print(f"[DIAGNOSTIC] get_discovery_prompt: Mode={mode}, ContentLength={len(content)}")
    if mode == "SD":
        schema_info = """
        - rd: Sale Date
        - amount: Figures
        - amount_words: Words
        - hypothecation: Bank Name if property is mortgaged
        - ss[i]: Sellers. n=Name, a=Age, c=Caste, relation_text=Relation details (e.g. S/o Late Mr. X or Wife of Mr. Y), adr=Address, id=Aadhar, pan=PAN
        - bs[i]: Buyers. n=Name, a=Age, c=Caste, relation_text=Relation details, adr=Address, id=Aadhar, pan=PAN
        - ws[i]: Witnesses. n=Name, relation_text=Relation details, adr=Address
        - ps[0]: Property details. plot_no, scheme, village, tehsil, dist, land_area, const_area, unit, adr, n, s, e, w
        - title_chain[i]: Title chain transitions. owner, deed_type, date, book, vol, page, reg_no, add_book
        - reg: {office, book, vol, page, reg_no, reg_date}
        """
    else:
        schema_info = """
        - rd: RM Execution Date
        - ad: Loan Agreement Date
        - bs[i]: Borrowers. s=Salutation, n=Name, a=Age, r=Relation, rn=Rel Name, adr=Address, id=Aadhar/ID, pan=PAN Card
        - ls[i]: Loans. n=LAN No, a=Amount in Figures, w=Amount in Words, t=Tenure
        - ps[i]: Properties. adr=Address, n=North, s=South, e=East, w=West
        - bsign: Bank Signatory. n=Name, a=Age, r=Relation, rn=Relative Name, id=Aadhar/ID, pan=PAN Card
        - ws[i]: Witnesses. n=Name, r=Relation, rn=Relative Name, adr=Address
        """

    return f"""
    CRITICAL MISSION: CONVERT COMPLETED DOCUMENT TO MASTER JINJA2 TEMPLATE.
    Identify ALL case-specific variable fields in the text below and map them to our system tags.

    MODE: {mode}
    CORE TAG SCHEMA:
    {schema_info}

    STRICT CONSTRAINTS:
    1. Return ONLY a JSON dictionary where keys are EXACT text from the document and values are the tags.
    2. Example: {{"24th March 2026": "{{{{rd}}}}", "15,00,000": "{{{{ls[0].a}}}}"}}
    3. DO NOT include any keys mapped to null, empty string, or any other value. ONLY include keys that are successfully mapped to our Jinja2 tags.
    4. Ensure each unique document string appears as a key exactly once. DO NOT duplicate keys or repeat mappings.
    5. STRICT RULE: DO NOT INCLUDE ANY CONVERSATIONAL TEXT, PREAMBLES, OR EXPLANATIONS. ONLY THE JSON OBJECT.
    6. SEPARATE NAMES AND RELATION DETAILS: If a name appears with a relation string in the document (e.g. "Jh foLoukFk iq= Jh Hkxoku flag" or "Jh foosd lDlSuk iq= Lo- Jh ts-ch- lDlSuk"), you MUST map them separately. Map the name part (e.g., "Jh foLoukFk") to the name tag (e.g., "{{ss[0].n}}") and the relation part (e.g., "iq= Jh Hkxoku flag" or "iq= Lo- Jh ts-ch- lDlSuk") to the relation_text tag (e.g., "{{ss[0].relation_text}}"). DO NOT map the entire combined string to the name tag!
    7. OCR RULE: OCR content must never be replaced by witness, seller, or buyer placeholders. OCR remains isolated.

    TEXT:
    {content}
    """
