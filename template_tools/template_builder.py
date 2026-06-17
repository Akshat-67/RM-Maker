import os
import sys
import json
import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from docx import Document
from docx.document import Document as _Document
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.text.run import Run
from copy import deepcopy

# Add root directory to path so extractor can be found
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

DEFAULT_API_KEYS = [
    os.getenv("GEMINI_API_KEY"),
    "AIzaSyBKkvXn_tTX2YVK_YzQPkm7FUDtYju43hc",
    "AIzaSyAXF1GYok40JQPkzg3rv2b_CGVJjDsaze8",
    "AQ.Ab8RN6LT45vwhXCygf42E1_cCExGjyFvD95qNo1vQ37hC3ZnBQ",
    "AIzaSyDs32YIJx35FDhb9qOa3vTcWDtU-RpL5_w"
]
DEFAULT_API_KEYS = list(dict.fromkeys([k for k in DEFAULT_API_KEYS if k]))
if not DEFAULT_API_KEYS:
    DEFAULT_API_KEYS = ["AIzaSyDs32YIJx35FDhb9qOa3vTcWDtU-RpL5_w"]

class TemplateBuilder:
    def __init__(self, root):
        self.root = root
        self.root.title("LegalDoc Master Template Creator v5")
        self.root.geometry("1450x980")

    def get_extractor(self, api_key=None, api_keys=None):
        if self.doc_type.get() == "SD":
            from modules.sd.extractor import SDDataExtractor as ExtractorClass
        else:
            from modules.rm.extractor import RMDataExtractor as ExtractorClass
        return ExtractorClass(api_key=api_key, api_keys=api_keys)

        self.doc_path = ""
        self.mapping = {}
        self.checklist_items = []
        self.available_models = []
        
        # Mode selector
        self.doc_type = tk.StringVar(value="RM")
        
        # Field catalogs
        self.rm_fields = [
            ("RM Execution Date", "{{rd}}"),
            ("Loan Agreement Date", "{{ad}}"),
            
            ("Borrower 1 - Salutation", "{{bs[0].s}}"),
            ("Borrower 1 - Name", "{{bs[0].n}}"),
            ("Borrower 1 - Age", "{{bs[0].a}}"),
            ("Borrower 1 - Relation", "{{bs[0].r}}"),
            ("Borrower 1 - Rel Name", "{{bs[0].rn}}"),
            ("Borrower 1 - Address", "{{bs[0].adr}}"),
            ("Borrower 1 - Aadhar/ID", "{{bs[0].id}}"),
            ("Borrower 1 - PAN Card", "{{bs[0].pan}}"),
            
            ("Borrower 2 - Salutation", "{{bs[1].s}}"),
            ("Borrower 2 - Name", "{{bs[1].n}}"),
            ("Borrower 2 - Age", "{{bs[1].a}}"),
            ("Borrower 2 - Relation", "{{bs[1].r}}"),
            ("Borrower 2 - Rel Name", "{{bs[1].rn}}"),
            ("Borrower 2 - Address", "{{bs[1].adr}}"),
            ("Borrower 2 - Aadhar/ID", "{{bs[1].id}}"),
            ("Borrower 2 - PAN Card", "{{bs[1].pan}}"),
            
            ("Loan 1 - LAN No", "{{ls[0].n}}"),
            ("Loan 1 - Amount (Figures)", "{{ls[0].a}}"),
            ("Loan 1 - Amount (Words)", "{{ls[0].w}}"),
            ("Loan 1 - Tenure", "{{ls[0].t}}"),
            
            ("Loan 2 - LAN No", "{{ls[1].n}}"),
            ("Loan 2 - Amount (Figures)", "{{ls[1].a}}"),
            ("Loan 2 - Amount (Words)", "{{ls[1].w}}"),
            ("Loan 2 - Tenure", "{{ls[1].t}}"),
            
            ("Property 1 - Address", "{{ps[0].adr}}"),
            ("Property 1 - North", "{{ps[0].n}}"),
            ("Property 1 - South", "{{ps[0].s}}"),
            ("Property 1 - East", "{{ps[0].e}}"),
            ("Property 1 - West", "{{ps[0].w}}"),
            
            ("Property 2 - Address", "{{ps[1].adr}}"),
            ("Property 2 - North", "{{ps[1].n}}"),
            ("Property 2 - South", "{{ps[1].s}}"),
            ("Property 2 - East", "{{ps[1].e}}"),
            ("Property 2 - West", "{{ps[1].w}}"),
            
            ("Bank Signatory - Name", "{{bsign.n}}"),
            ("Bank Signatory - Age", "{{bsign.a}}"),
            ("Bank Signatory - Relation", "{{bsign.r}}"),
            ("Bank Signatory - Rel Name", "{{bsign.rn}}"),
            ("Bank Signatory - PAN Card", "{{bsign.pan}}"),
            ("Bank Signatory - Aadhar/ID", "{{bsign.id}}"),
            
            ("Witness 1 - Name", "{{ws[0].n}}"),
            ("Witness 1 - Relation", "{{ws[0].r}}"),
            ("Witness 1 - Rel Name", "{{ws[0].rn}}"),
            ("Witness 1 - Address", "{{ws[0].adr}}"),
            
            ("Witness 2 - Name", "{{ws[1].n}}"),
            ("Witness 2 - Relation", "{{ws[1].r}}"),
            ("Witness 2 - Rel Name", "{{ws[1].rn}}"),
            ("Witness 2 - Address", "{{ws[1].adr}}"),
            
            ("Document Schedule / Title Chain (FIRST SCHEDULE)", "{{ds_text}}"),
            ("Document Schedule 1 - Title Deed", "{{ds[0].t}}"),
            ("Second Schedule (Documents to be collected)", "{{second_schedule}}"),
        ]
        
        self.sd_fields = [
            ("Sale Date", "{{rd}}"),
            ("Total Sale Amount", "{{amount}}"),
            ("Sale Amount in Words", "{{amount_words}}"),
            ("Hypothecation (Bank)", "{{hypothecation}}"),
            
            ("Seller 1 - Name", "{{ss[0].n}}"),
            ("Seller 1 - Age", "{{ss[0].a}}"),
            ("Seller 1 - Caste", "{{ss[0].c}}"),
            ("Seller 1 - Relation", "{{ss[0].r}}"),
            ("Seller 1 - Rel Name", "{{ss[0].rn}}"),
            ("Seller 1 - Address", "{{ss[0].adr}}"),
            ("Seller 1 - Aadhar", "{{ss[0].id}}"),
            ("Seller 1 - PAN", "{{ss[0].pan}}"),
            
            ("Buyer 1 - Name", "{{bs[0].n}}"),
            ("Buyer 1 - Age", "{{bs[0].a}}"),
            ("Buyer 1 - Caste", "{{bs[0].c}}"),
            ("Buyer 1 - Relation", "{{bs[0].r}}"),
            ("Buyer 1 - Rel Name", "{{bs[0].rn}}"),
            ("Buyer 1 - Address", "{{bs[0].adr}}"),
            ("Buyer 1 - Aadhar", "{{bs[0].id}}"),
            ("Buyer 1 - PAN", "{{bs[0].pan}}"),
            
            ("Property - Plot No", "{{ps[0].plot_no}}"),
            ("Property - Scheme", "{{ps[0].scheme}}"),
            ("Property - Village", "{{ps[0].village}}"),
            ("Property - Tehsil", "{{ps[0].tehsil}}"),
            ("Property - District", "{{ps[0].dist}}"),
            ("Property - Land Area", "{{ps[0].land_area}}"),
            ("Property - Const Area", "{{ps[0].const_area}}"),
            ("Property - Unit", "{{ps[0].unit}}"),
            ("Property - Address", "{{ps[0].adr}}"),
            ("Property - North", "{{ps[0].n}}"),
            ("Property - South", "{{ps[0].s}}"),
            ("Property - East", "{{ps[0].e}}"),
            ("Property - West", "{{ps[0].w}}"),
            
            ("Witness 1 - Name", "{{ws[0].n}}"),
            ("Witness 1 - Relation", "{{ws[0].r}}"),
            ("Witness 1 - Rel Name", "{{ws[0].rn}}"),
            ("Witness 1 - Address", "{{ws[0].adr}}"),
            
            ("Title 1 - Owner", "{{title_chain[0].owner}}"),
            ("Title 1 - Deed Type", "{{title_chain[0].deed_type}}"),
            ("Title 1 - Date", "{{title_chain[0].date}}"),
            ("Title 1 - Book", "{{title_chain[0].book}}"),
            ("Title 1 - Vol", "{{title_chain[0].vol}}"),
            ("Title 1 - Page", "{{title_chain[0].page}}"),
            ("Title 1 - Reg No", "{{title_chain[0].reg_no}}"),
            ("Title 1 - Add Book", "{{title_chain[0].add_book}}"),
            
            ("Reg - Office", "{{reg.office}}"),
            ("Reg - Book", "{{reg.book}}"),
            ("Reg - Vol", "{{reg.vol}}"),
            ("Reg - Page", "{{reg.page}}"),
            ("Reg - No", "{{reg.reg_no}}"),
            ("Reg - Date", "{{reg.reg_date}}"),
        ]
        
        self.expected_fields = self.rm_fields
        self.setup_ui()
        # Auto-fetch models in background
        threading.Thread(target=self._auto_fetch_models, daemon=True).start()

    def _auto_fetch_models(self):
        """Fetch available Gemini models on startup."""
        try:
            extractor = self.get_extractor(api_keys=DEFAULT_API_KEYS)
            models = extractor.get_available_models()
            if isinstance(models, list) and models:
                self.available_models = models
                preferred = [m for m in models if "flash" in m.lower()]
                best = preferred[0] if preferred else models[0]
                if self.model_var.get() not in ["NVIDIA Nemotron-3 550B", "NVIDIA Llama 3.1 8B"]:
                    self.model_var.set(best)
                self.root.after(0, self._update_model_dropdown)
        except Exception:
            pass

    def _update_model_dropdown(self):
        if hasattr(self, 'model_dropdown') and self.model_dropdown:
            nvidia_models = ["NVIDIA Nemotron-3 550B", "NVIDIA Llama 3.1 8B"]
            models_list = nvidia_models + [m for m in self.available_models if m not in nvidia_models]
            self.model_dropdown['values'] = models_list
            if self.model_var.get() not in models_list:
                self.model_var.set("NVIDIA Nemotron-3 550B")

    def rotate_api_key_click(self):
        if not hasattr(self, 'current_key_idx'):
            self.current_key_idx = 0
        self.current_key_idx = (self.current_key_idx + 1) % len(DEFAULT_API_KEYS)
        new_key = DEFAULT_API_KEYS[self.current_key_idx]
        self.api_key_entry.delete(0, tk.END)
        self.api_key_entry.insert(0, new_key)
        # Fetch models with the new key in the background
        threading.Thread(target=self._refresh_models_with_key, args=(new_key,), daemon=True).start()
        messagebox.showinfo("API Key Switched", f"Switched to API Key #{self.current_key_idx + 1} (ends with ...{new_key[-4:]})")

    def _refresh_models_with_key(self, api_key):
        try:
            extractor = self.get_extractor(api_key=api_key)
            models = extractor.get_available_models()
            if isinstance(models, list) and models:
                self.available_models = models
                preferred = [m for m in models if "flash" in m.lower()]
                best = preferred[0] if preferred else models[0]
                if self.model_var.get() not in ["NVIDIA Nemotron-3 550B", "NVIDIA Llama 3.1 8B"]:
                    self.model_var.set(best)
                self.root.after(0, self._update_model_dropdown)
        except Exception:
            pass

    def setup_ui(self):
        self.c_bg = "#F3F4F6"; self.c_panel = "#FFFFFF"; self.c_blue = "#1A73E8"; self.c_success = "#0F9D58"; self.c_danger = "#D93025"; self.c_border = "#DADCE0"; self.c_text = "#3C4043"
        self.root.configure(bg=self.c_bg)
        
        header = tk.Frame(self.root, bg=self.c_blue, height=70)
        header.pack(fill="x")
        tk.Label(header, text="MASTER TEMPLATE GENERATOR PRO", fg="white", bg=self.c_blue, font=("Segoe UI", 18, "bold"), padx=25).pack(side="left")
        
        main_content = tk.Frame(self.root, bg=self.c_bg)
        main_content.pack(fill="both", expand=True, padx=20, pady=15)
        
        # Left Panel
        left_panel = tk.Frame(main_content, bg=self.c_bg, width=520)
        left_panel.pack(side="left", fill="y", padx=(0, 15))
        left_panel.pack_propagate(False)
        
        # 1. Select Source
        group1 = tk.LabelFrame(left_panel, text=" 1. SELECT SOURCE WORD DOCUMENT ", bg=self.c_panel, font=("Segoe UI", 10, "bold"), padx=15, pady=15, fg=self.c_blue)
        group1.pack(fill="x", pady=(0, 10))
        
        mode_frame = tk.Frame(group1, bg=self.c_panel)
        mode_frame.pack(fill="x", pady=(0, 10))
        tk.Label(mode_frame, text="Document Mode:", bg=self.c_panel, font=("Segoe UI", 9, "bold")).pack(side="left")
        tk.Radiobutton(mode_frame, text="RM", variable=self.doc_type, value="RM", command=self.switch_mode, bg=self.c_panel).pack(side="left", padx=10)
        tk.Radiobutton(mode_frame, text="SD", variable=self.doc_type, value="SD", command=self.switch_mode, bg=self.c_panel).pack(side="left")

        tk.Button(group1, text="OPEN .DOCX FILE", command=self.load_doc, bg=self.c_blue, fg="white", bd=0, pady=12, font=("Segoe UI", 10, "bold"), cursor="hand2").pack(fill="x")
        self.path_lbl = tk.Label(group1, text="No file selected", fg="#70757A", bg=self.c_panel, font=("Segoe UI", 9))
        self.path_lbl.pack(pady=8)
        
        # 2. AI Config
        group2 = tk.LabelFrame(left_panel, text=" 2. AI CONFIGURATION ", bg=self.c_panel, font=("Segoe UI", 10, "bold"), padx=15, pady=15, fg=self.c_blue)
        group2.pack(fill="x", pady=10)
        
        tk.Label(group2, text="API Key:", bg=self.c_panel, font=("Segoe UI", 9)).pack(anchor="w")
        self.api_key_entry = tk.Entry(group2, show="*", bg="#F1F3F4", bd=0, font=("Consolas", 10))
        self.api_key_entry.insert(0, DEFAULT_API_KEYS[0] if DEFAULT_API_KEYS else "")
        self.api_key_entry.pack(fill="x", ipady=8, pady=5)
        
        tk.Button(group2, text="🔄 Rotate / Switch API Key", command=self.rotate_api_key_click, bg="#E8F0FE", fg=self.c_blue, bd=0, font=("Segoe UI", 8, "bold"), cursor="hand2").pack(fill="x", pady=2)
        
        tk.Label(group2, text="AI Model:", bg=self.c_panel, font=("Segoe UI", 9)).pack(anchor="w", pady=(5,0))
        self.model_var = tk.StringVar(value="NVIDIA Nemotron-3 550B")
        self.model_dropdown = ttk.Combobox(group2, textvariable=self.model_var, values=["NVIDIA Nemotron-3 550B", "NVIDIA Llama 3.1 8B", "gemini-2.5-flash", "gemini-1.5-flash"], state="readonly", font=("Segoe UI", 10))
        self.model_dropdown.pack(fill="x", pady=5)
        tk.Button(group2, text="Refresh Models", command=self.refresh_models, bg="#E8F0FE", fg=self.c_blue, bd=0, font=("Segoe UI", 8, "bold")).pack(fill="x", pady=5)
        
        # Discovery Button
        self.discover_btn = tk.Button(left_panel, text="🔍 START FULL AI DATA DISCOVERY", command=self.discover_data, bg=self.c_success, fg="white", font=("Segoe UI", 13, "bold"), bd=0, pady=18, cursor="hand2")
        self.discover_btn.pack(fill="x", pady=15)
        
        # 3. Discovered Mappings
        group3 = tk.LabelFrame(left_panel, text=" 3. CURRENT MAPPING AUDIT ", bg=self.c_panel, font=("Segoe UI", 10, "bold"), padx=10, pady=10, fg=self.c_blue)
        group3.pack(fill="both", expand=True, pady=10)
        
        self.tree = ttk.Treeview(group3, columns=("Raw", "Tag"), show='headings')
        self.tree.heading("Raw", text="Exact Document Substring")
        self.tree.heading("Tag", text="Placeholder Tag (Jinja2)")
        self.tree.column("Raw", width=320)
        self.tree.column("Tag", width=160)
        self.tree.pack(side="left", fill="both", expand=True)
        
        sb = ttk.Scrollbar(group3, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        
        # Right Panel
        right_panel = tk.Frame(main_content, bg=self.c_bg, width=600)
        right_panel.pack(side="left", fill="both", expand=True, padx=(10, 0))
        
        # Smart Gap Fixer / Manual Mapping
        r1 = tk.LabelFrame(right_panel, text=" SMART MANUAL MAPPING ", bg=self.c_panel, font=("Segoe UI", 10, "bold"), padx=15, pady=15, fg=self.c_text)
        r1.pack(fill="x", pady=(0, 10))
        
        tk.Label(r1, text="Step 1: Type the EXACT text from the document:", bg=self.c_panel, font=("Segoe UI", 9), anchor="w").pack(fill="x")
        self.manual_key = tk.Entry(r1, bg="#F1F3F4", bd=0, font=("Segoe UI", 10), highlightthickness=1, highlightbackground=self.c_border)
        self.manual_key.pack(fill="x", ipady=10, pady=5)
        
        tk.Label(r1, text="Step 2: Choose the placeholder tag from the dropdown:", bg=self.c_panel, font=("Segoe UI", 9), anchor="w").pack(fill="x")
        self.manual_tag_display = tk.StringVar()
        tag_display_values = [f"{label}  →  {tag}" for label, tag in self.expected_fields]
        self.manual_tag_dropdown = ttk.Combobox(r1, textvariable=self.manual_tag_display, values=tag_display_values, state="normal", font=("Segoe UI", 10))
        self.manual_tag_dropdown.pack(fill="x", ipady=4, pady=5)
        
        btn_frame = tk.Frame(r1, bg=self.c_panel)
        btn_frame.pack(fill="x", pady=5)
        tk.Button(btn_frame, text="+ ADD MAPPING", command=self.add_manual, bg=self.c_blue, fg="white", bd=0, pady=10, font=("Segoe UI", 9, "bold"), cursor="hand2").pack(side="left", padx=(0, 10))
        tk.Button(btn_frame, text="REMOVE SELECTED", command=self.delete_selected, fg=self.c_danger, bg=self.c_panel, bd=1, pady=10, cursor="hand2").pack(side="left")
        
        # Verification Checklist
        ver_label = tk.Label(right_panel, text="VERIFICATION CHECKLIST:", bg=self.c_bg, fg=self.c_text, font=("Segoe UI", 9, "bold"), anchor="w")
        ver_label.pack(fill="x", pady=(10, 2))
        
        ver_frame_wrapper = tk.Frame(right_panel, bg=self.c_panel, highlightthickness=1, highlightbackground=self.c_border)
        ver_frame_wrapper.pack(fill="both", expand=True)
        
        ver_canvas = tk.Canvas(ver_frame_wrapper, bg=self.c_panel, highlightthickness=0)
        ver_sb = ttk.Scrollbar(ver_frame_wrapper, orient="vertical", command=ver_canvas.yview)
        self.check_scroll_f = tk.Frame(ver_canvas, bg=self.c_panel)
        self.check_scroll_f.bind("<Configure>", lambda e: ver_canvas.configure(scrollregion=ver_canvas.bbox("all")))
        ver_canvas.create_window((0, 0), window=self.check_scroll_f, anchor="nw", width=560)
        ver_canvas.configure(yscrollcommand=ver_sb.set)
        ver_canvas.pack(side="left", fill="both", expand=True)
        ver_sb.pack(side="right", fill="y")
        self.init_checklist()
        
        # Save Button
        self.save_btn = tk.Button(right_panel, text="GENERATE FINAL MASTER TEMPLATE", command=self.generate_master, bg=self.c_blue, fg="white", font=("Segoe UI", 14, "bold"), bd=0, pady=25, cursor="hand2")
        self.save_btn.pack(fill="x", pady=(15, 5))

    def init_checklist(self):
        for w in self.check_scroll_f.winfo_children(): w.destroy()
        self.checklist_items = []
        for label, tag in self.expected_fields:
            f = tk.Frame(self.check_scroll_f, bg=self.c_panel)
            f.pack(fill="x", pady=4)
            
            var = tk.BooleanVar(value=False)
            cb = tk.Checkbutton(f, text=f"{label}  ({tag})", variable=var, bg=self.c_panel, font=("Segoe UI", 10))
            cb.pack(side="left")
            self.checklist_items.append((label, tag, var, cb))

    def switch_mode(self):
        if self.doc_type.get() == "SD":
            self.expected_fields = self.sd_fields
        else:
            self.expected_fields = self.rm_fields
        
        # Update dropdown values
        tag_display_values = [f"{label}  →  {tag}" for label, tag in self.expected_fields]
        self.manual_tag_dropdown['values'] = tag_display_values
        self.manual_tag_display.set("")
        
        # Refresh checklist
        self.init_checklist()
        self.update_checklist_status()

    def update_checklist_status(self):
        all_mapped_tags = " ".join(self.mapping.values())
        for label, tag, var, cb in self.checklist_items:
            if tag in all_mapped_tags:
                var.set(True)
                cb.config(fg=self.c_success)
            else:
                var.set(False)
                cb.config(fg=self.c_danger)

    def load_doc(self):
        p = filedialog.askopenfilename(filetypes=[("Word Document", "*.docx")])
        if p:
            self.doc_path = p
            self.path_lbl.config(text=os.path.basename(p), fg=self.c_blue)

    def refresh_models(self):
        k = self.api_key_entry.get().strip()
        if not k:
            messagebox.showwarning("Key Required", "Provide Gemini API Key.")
            return
        try:
            extractor = self.get_extractor(api_key=k)
            models = extractor.get_available_models()
            if isinstance(models, list) and models:
                self.available_models = models
                nvidia_models = ["NVIDIA Nemotron-3 550B", "NVIDIA Llama 3.1 8B"]
                models_list = nvidia_models + [m for m in models if m not in nvidia_models]
                self.model_dropdown['values'] = models_list
                if self.model_var.get() not in models_list:
                    self.model_var.set("NVIDIA Nemotron-3 550B")
                messagebox.showinfo("AI System", f"Successfully found {len(models)} models.")
            else:
                messagebox.showwarning("AI System", "No models found or error returned.")
        except Exception as e:
            messagebox.showerror("AI Error", str(e))

    def discover_data(self):
        if not self.doc_path:
            messagebox.showwarning("Warning", "Load a document first!")
            return
        if not self.api_key_entry.get():
            messagebox.showwarning("Warning", "Provide an API key!")
            return
            
        self.discover_btn.config(state="disabled", text="AI DISCOVERY IN PROGRESS...")
        threading.Thread(target=self.run_ai_discovery, daemon=True).start()

    def run_ai_discovery(self):
        try:
            doc = Document(self.doc_path)
            content = self.get_doc_content(doc)
            
            mode = self.doc_type.get()
            if mode == "SD":
                schema_info = """
                - rd: Sale Date
                - amount: Figures
                - amount_words: Words
                - ss[i]: Sellers. n=Name, a=Age, c=Caste, relation_text=Relation details (e.g. S/o Late Mr. X or Wife of Mr. Y), adr=Address, id=Aadhar, pan=PAN
                - bs[i]: Buyers. n=Name, a=Age, c=Caste, relation_text=Relation details, adr=Address, id=Aadhar, pan=PAN
                - ps[0]: Property. plot_no, scheme, village, tehsil, dist, land_area, const_area, unit, adr=Address, n=North, s=South, e=East, w=West
                - ws[i]: Witnesses. n=Name, relation_text=Relation details, adr=Address
                - title_chain[i]: {owner, deed_type, date, book, vol, page, reg_no, add_book}
                - reg: {office, book, vol, page, reg_no, reg_date}
                - hypothecation: Bank Name if property is mortgaged
                """
            else:
                schema_info = """
                - rd: RM Execution Date
                - ad: Loan Agreement Date
                - bs[i]: Borrowers. s=Salutation, n=Name, a=Age, r=Relation, rn=Rel Name, adr=Address, id=Aadhar/ID, pan=PAN Card
                - ls[i]: Loans. n=LAN No, a=Amount in Figures, w=Amount in Words, t=Tenure
                - ps[i]: Properties. adr=Address, n=North, s=South, e=East, w=West
                - bsign: Bank Signatory. n=Name, a=Age, r=Relation, rn=Relative Name, id=Aadhar/ID, pan=PAN Card
                - ws[i]: Witnesses. n=Name, r=Relation, rn=Relative Name, adr=Address
                """

            prompt = f"""
            CRITICAL MISSION: CONVERT COMPLETED DOCUMENT TO MASTER JINJA2 TEMPLATE.
            Identify ALL case-specific variable fields in the text below and map them to our system tags.

            MODE: {mode}
            CORE TAG SCHEMA:
            {schema_info}

            STRICT CONSTRAINTS:
            1. Return ONLY a JSON dictionary where keys are EXACT text from the document and values are the tags.
            2. Example: {{"24th March 2026": "{{{{rd}}}}", "15,00,000": "{{{{ls[0].a}}}}"}}
            3. DO NOT include any keys mapped to null, empty string, or any other value. ONLY include keys that are successfully mapped to our Jinja2 tags.
            4. Ensure each unique document string appears as a key exactly once. DO NOT duplicate keys or repeat mappings.
            5. STRICT RULE: DO NOT INCLUDE ANY CONVERSATIONAL TEXT, PREAMBLES, OR EXPLANATIONS. ONLY THE JSON OBJECT.
            6. SEPARATE NAMES AND RELATION DETAILS: If a name appears with a relation string in the document (e.g. "Jh foLoukFk iq= Jh Hkxoku flag" or "Jh foosd lDlSuk iq= Lo- Jh ts-ch- lDlSuk"), you MUST map them separately. Map the name part (e.g., "Jh foLoukFk") to the name tag (e.g., "{{ss[0].n}}") and the relation part (e.g., "iq= Jh Hkxoku flag" or "iq= Lo- Jh ts-ch- lDlSuk") to the relation_text tag (e.g., "{{ss[0].relation_text}}"). DO NOT map the entire combined string to the name tag!
            7. OCR RULE: OCR content must never be replaced by witness, seller, or buyer placeholders. OCR remains isolated.

            TEXT:
            {content}
            """
            
            extractor = self.get_extractor(api_key=self.api_key_entry.get().strip())
            raw = extractor.raw_generate(prompt, self.model_var.get())
            
            match = re.search(r'\{.*\}', raw or "", re.DOTALL)
            if match:
                mapping_str = match.group(0)
                # Escape double quotes inside DevLys JSON keys and values
                cleaned_lines = []
                for line in mapping_str.splitlines():
                    line_stripped = line.strip()
                    if line_stripped in ('{', '}', '[', ']', '') or not line_stripped.startswith('"'):
                        cleaned_lines.append(line)
                        continue
                    match_val = re.search(r':\s*(("(.*)"|null|true|false|\d+)\s*(,?)\s*)$', line_stripped)
                    if not match_val:
                        cleaned_lines.append(line)
                        continue
                    key_part = line_stripped[:match_val.start()].strip()
                    val_part = match_val.group(1).strip()
                    if key_part.startswith('"') and key_part.endswith('"'):
                        raw_key = key_part[1:-1]
                        escaped_key = raw_key.replace('\\"', '"').replace('"', '\\"')
                        key_part = f'"{escaped_key}"'
                    raw_val_match = re.match(r'^"(.*)"(,?)$', val_part)
                    if raw_val_match:
                        raw_val = raw_val_match.group(1)
                        comma = raw_val_match.group(2)
                        escaped_val = raw_val.replace('\\"', '"').replace('"', '\\"')
                        val_part = f'"{escaped_val}"{comma}'
                    indent = line[:len(line) - len(line.lstrip())]
                    cleaned_lines.append(f'{indent}{key_part}: {val_part}')
                cleaned_mapping_str = "\n".join(cleaned_lines)
                
                try:
                    mapping = json.loads(cleaned_mapping_str)
                except Exception:
                    mapping = json.loads(mapping_str)
                    
                self.mapping = self.clean_mapping(mapping)
                self.root.after(0, self.update_tree)
            else:
                self.root.after(0, lambda: messagebox.showwarning("AI Scan", "Gemini could not identify structured fields."))
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Error", str(e)))
        finally:
            self.root.after(0, lambda: self.discover_btn.config(state="normal", text="🔍 START FULL AI DATA DISCOVERY"))

    def get_doc_content(self, doc):
        chunks = []
        for part_name, part in self.iter_story_parts(doc):
            for p in self.iter_paragraphs_deep(part):
                if p.text.strip():
                    chunks.append(p.text.strip())
        return "\n".join(chunks)

    def iter_block_items(self, parent):
        if isinstance(parent, _Document):
            parent_elm = parent.element.body
        elif isinstance(parent, _Cell):
            parent_elm = parent._tc
        else:
            parent_elm = parent._element

        for child in parent_elm.iterchildren():
            if child.tag.endswith('}p'):
                yield Paragraph(child, parent)
            elif child.tag.endswith('}tbl'):
                yield Table(child, parent)

    def iter_paragraphs_deep(self, parent):
        for block in self.iter_block_items(parent):
            if isinstance(block, Paragraph):
                yield block
            elif isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        yield from self.iter_paragraphs_deep(cell)

    def iter_story_parts(self, doc):
        yield "body", doc
        for section in doc.sections:
            for hf_attr in ['header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer']:
                hf = getattr(section, hf_attr, None)
                if hf:
                    yield hf_attr, hf

    def clean_mapping(self, mapping):
        if not isinstance(mapping, dict): return {}
        clean = {}
        for k, v in mapping.items():
            if k and v and "{{" in str(v):
                clean[str(k).strip()] = str(v).strip()
        return clean

    def update_tree(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        for k, v in sorted(self.mapping.items(), key=lambda x: len(x[0]), reverse=True):
            self.tree.insert("", "end", values=(k, v))
        self.update_checklist_status()

    def add_manual(self):
        k = self.manual_key.get().strip()
        v_display = self.manual_tag_display.get().strip()
        if "→" in v_display:
            v = v_display.split("→")[1].strip()
        else:
            v = v_display
            
        if k and v and "{{" in v:
            self.mapping[k] = v
            self.update_tree()
            self.manual_key.delete(0, tk.END)

    def delete_selected(self):
        selected = self.tree.selection()
        for i in selected:
            values = self.tree.item(i)['values']
            if values:
                k = values[0]
                if k in self.mapping:
                    del self.mapping[k]
        self.update_tree()

    def generate_master(self):
        if not self.doc_path or not self.mapping:
            messagebox.showwarning("Warning", "No document or mappings!")
            return
            
        doc = Document(self.doc_path)
        reps = sorted(self.mapping.items(), key=lambda x: len(x[0]), reverse=True)
        
        for part_name, part in self.iter_story_parts(doc):
            for p in self.iter_paragraphs_deep(part):
                for old, new in reps:
                    self.safe_replace(p, old, new)
        
        save_p = filedialog.asksaveasfilename(defaultextension=".docx", initialfile="MASTER_TEMPLATE_READY.docx")
        if save_p:
            doc.save(save_p)
            messagebox.showinfo("Done", "Master Template generated successfully!")

    def safe_replace(self, p, old, new):
        if not old or old == new: return
        match = self.find_fuzzy_match(p, old)
        if match:
            s_r, s_o, e_r, e_o = match
            is_tag = "{{" in new
            
            if s_r == e_r:
                run = p.runs[s_r]
                if is_tag:
                    # Surgical run splitting to preserve formatting but force readable font for tag
                    from copy import deepcopy
                    from docx.text.run import Run
                    
                    orig_text = run.text
                    prefix = orig_text[:s_o]
                    suffix = orig_text[e_o:]
                    
                    # Current run becomes prefix
                    run.text = prefix
                    
                    # Create tag run
                    tag_el = deepcopy(run._element)
                    run._element.addnext(tag_el)
                    tag_run = Run(tag_el, run._parent)
                    tag_run.text = new
                    tag_run.font.name = "Arial"
                    # Force Arial across all font slots to avoid DevLys rendering
                    tag_run._element.rPr.get_or_add_rFonts().set(qn('w:ascii'), 'Arial')
                    tag_run._element.rPr.get_or_add_rFonts().set(qn('w:hAnsi'), 'Arial')
                    tag_run._element.rPr.get_or_add_rFonts().set(qn('w:cs'), 'Arial')
                    
                    # Create suffix run
                    suffix_el = deepcopy(run._element)
                    tag_el.addnext(suffix_el)
                    suffix_run = Run(suffix_el, run._parent)
                    suffix_run.text = suffix
                else:
                    run.text = run.text[:s_o] + new + run.text[e_o:]
            else:
                # Multi-run replacement
                p.runs[s_r].text = p.runs[s_r].text[:s_o] + new
                for i in range(s_r + 1, e_r):
                    p.runs[i].text = ""
                p.runs[e_r].text = p.runs[e_r].text[e_o:]
                
                if is_tag:
                    p.runs[s_r].font.name = "Arial"
                    try:
                        p.runs[s_r]._element.rPr.get_or_add_rFonts().set(qn('w:ascii'), 'Arial')
                        p.runs[s_r]._element.rPr.get_or_add_rFonts().set(qn('w:hAnsi'), 'Arial')
                    except: pass

    def find_fuzzy_match(self, p, needle):
        if not p.runs: return None
        chars = []
        pos = []
        for ri, run in enumerate(p.runs):
            for ci, c in enumerate(run.text):
                chars.append(" " if c.isspace() or c == "\u00A0" else c)
                pos.append((ri, ci))
        
        haystack = "".join(chars)
        target = re.sub(r"[\s\u00A0]+", " ", needle).strip().casefold()
        pattern = re.escape(target).replace(r"\ ", r"[\s\u00A0]+")
        
        try:
            m = re.search(pattern, haystack, re.IGNORECASE)
            if m:
                s = m.start()
                e = m.end()
                sr, so = pos[s]
                er, eo = pos[e-1]
                return sr, so, er, eo + 1
        except:
            pass
        return None

if __name__ == "__main__":
    root = tk.Tk()
    TemplateBuilder(root)
    root.mainloop()
