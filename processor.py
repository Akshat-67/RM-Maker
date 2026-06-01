from docxtpl import DocxTemplate, RichText
from docx.enum.text import WD_COLOR_INDEX
import os
import re
import zipfile
from html import unescape

# Legacy Hindi Font Converters (Unicode <=> DevLys 010/040 / KrutiDev 010)
def KrutiDev_to_Unicode(krutidev_substring):
    modified_substring = krutidev_substring
    array_one = ["ñ","Q+Z","sas","aa",")Z","ZZ","‘","’","“","”",
    "å",  "ƒ",  "„",   "…",   "†",   "‡",   "ˆ",   "‰",   "Š",   "‹", 
    "¶+",   "d+", "[+k","[+", "x+",  "T+",  "t+", "M+", "<+", "Q+", ";+", "j+", "u+",
    "Ùk", "Ù", "Dr", "–", "—","é","™","=kk","f=k",  
    "à",   "á",    "â",   "ã",   "ºz",  "º",   "í", "{k", "{", "=",  "«",   
    "Nî",   "Vî",    "Bî",   "Mî",   "<î", "|", "K", "}",
    "J",   "Vª",   "Mª",  "<ªª",  "Nª",   "Ø",  "Ý", "nzZ",  "æ", "ç", "Á", "xz", "#", ":",
    "v‚","vks",  "vkS",  "vk",    "v",  "b±", "Ã",  "bZ",  "b",  "m",  "Å",  ",s",  ",",   "_",
    "ô",  "d", "Dk", "D", "[k", "[", "x","Xk", "X", "Ä", "?k", "?",   "³", 
    "pkS",  "p", "Pk", "P",  "N",  "t", "Tk", "T",  ">", "÷", "¥",
    "ê",  "ë",   "V",  "B",   "ì",   "ï", "M+", "<+", "M",  "<", ".k", ".",    
    "r",  "Rk", "R",   "Fk", "F",  ")", "n", "/k", "èk",  "/", "Ë", "è", "u", "Uk", "U",   
    "i",  "Ik", "I",   "Q",    "¶",  "c", "Ck",  "C",  "Hk",  "H", "e", "Ek",  "E",
    ";",  "¸",   "j",    "y", "Yk",  "Y",  "G",  "o", "Ok", "O",
    "'k", "'",   "\"k",  "\"",  "l", "Lk",  "L",   "g", 
    "È", "z", 
    "Ì", "Í", "Î",  "Ï",  "Ñ",  "Ò",  "Ó",  "Ô",   "Ö",  "Ø",  "Ù","Ük", "Ü",
    "‚",    "ks",   "kS",   "k",  "h",    "q",   "w",   "`",    "s",    "S",
    "a",    "¡",    "%",     "W",  "•", "·", "∙", "·", "~j",  "~", "\\","+"," ः",
    "^", "*",  "Þ", "ß", "(", "¼", "½", "¿", "À", "¾", "A", "-", "&", "&", "Œ", "]","~ ","@"]
    
    array_two = ["॰","QZ+","sa","a","र्द्ध","Z","\"","\"","'","'",
    "०",  "१",  "२",  "३",     "४",   "५",  "६",   "७",   "८",   "९",   
    "फ़्",  "क़",  "ख़", "ख़्",  "ग़", "ज़्", "ज़",  "ड़",  "ढ़",   "फ़",  "य़",  "ऱ",  "ऩ",    
    "त्त", "त्त्", "क्त",  "दृ",  "कृ","न्न","न्न्","=k","f=",
    "ह्न",  "ह्य",  "हृ",  "ह्म",  "ह्र",  "ह्",   "द्द",  "क्ष", "क्ष्", "त्र", "त्र्", 
    "छ्य",  "ट्य",  "ठ्य",  "ड्य",  "ढ्य", "द्य", "ज्ञ", "द्व",
    "श्र",  "ट्र",    "ड्र",    "ढ्र",    "छ्र",   "क्र",  "फ्र", "र्द्र",  "द्र",   "प्र", "प्र",  "ग्र", "रु",  "रू",
    "ऑ",   "ओ",  "औ",  "आ",   "अ", "ईं", "ई",  "ई",   "इ",  "उ",   "ऊ",  "ऐ",  "ए", "ऋ",
    "क्क", "क", "क", "क्", "ख", "ख्", "ग", "ग", "ग्", "घ", "घ", "घ्", "ङ",
    "चै",  "च", "च", "च्", "छ", "ज", "ज", "ज्",  "झ",  "झ्", "ञ",
    "ट्ट",   "ट्ठ",   "ट",   "ठ",   "ड्ड",   "ड्ढ",  "ड़", "ढ़", "ड",   "ढ", "ण", "ण्",   
    "त", "त", "त्", "थ", "थ्",  "द्ध",  "द", "ध", "ध", "ध्", "ध्", "ध्", "न", "न", "न्",    
    "प", "प", "प्",  "फ", "फ्",  "ब", "ब", "ब्",  "भ", "भ्",  "म",  "म", "म्",  
    "य", "य्",  "र", "ल", "ल", "ल्",  "ळ",  "व", "व", "व्",   
    "श", "श्",  "ष", "ष्", "स", "स", "स्", "ह", 
    "ीं", "्र",    
    "द्द", "ट्ट","ट्ठ","ड्ड","कृ","भ","्य","ड्ढ","झ्","क्र","त्त्","श","श्",
    "ॉ",  "ो",   "ौ",   "ा",   "ी",   "ु",   "ू",   "ृ",   "े",   "ै",
    "ं",   "ँ",   "ः",   "ॅ",  "ऽ", "ऽ", "ऽ", "ऽ", "्र",  "्", "?", "़",":",
    "‘",   "’",   "“",   "”",  ";",  "(",    ")",   "{",    "}",   "=", "।", ".", "-",  "µ", "॰", ",","् ","/"]
    
    array_one_length = len(array_one)
    modified_substring = "  " + modified_substring + "  "
    position_of_f = modified_substring.rfind("f")
    while (position_of_f != -1):    
        modified_substring = modified_substring[:position_of_f] + modified_substring[position_of_f+1] + modified_substring[position_of_f] +  modified_substring[position_of_f+2:]
        position_of_f = modified_substring.rfind("f",0, position_of_f - 1 )
    modified_substring = modified_substring.replace("f","ि")
    modified_substring = modified_substring.strip()
    
    modified_substring = "  " + modified_substring + "  "
    position_of_r = modified_substring.find("Z")
    set_of_matras =  ["‚",    "ks",   "kS",   "k",     "h",    "q",   "w",   "`",    "s",    "S", "a",    "¡",    "%",     "W",   "·",   "~ ", "~"]
    while (position_of_r != -1):    
        modified_substring = modified_substring.replace("Z","",1)
        if modified_substring[position_of_r - 1] in set_of_matras:
            modified_substring = modified_substring[:position_of_r - 2] + "j~" + modified_substring[position_of_r - 2:]
        else:
            modified_substring = modified_substring[:position_of_r - 1] + "j~" + modified_substring[position_of_r - 1:]
        position_of_r = modified_substring.find("Z")
    modified_substring = modified_substring.strip()
    
    for input_symbol_idx in range(0, array_one_length):
        modified_substring = modified_substring.replace(array_one[input_symbol_idx ] , array_two[input_symbol_idx] )
    return modified_substring

def Unicode_to_KrutiDev(unicode_substring):
    modified_substring = unicode_substring
    array_one = ["‘",   "’",   "“",   "”",   "(",    ")",   "{",    "}",   "=", "।",  "?",  "-",  "µ", "॰", ",", ".", "् ", 
    "०",  "१",  "२",  "३",     "४",   "५",  "६",   "७",   "८",   "९", "x", 
    "फ़्",  "क़",  "ख़",  "ग़", "ज़्", "ज़",  "ड़",  "ढ़",   "फ़",  "य़",  "ऱ",  "ऩ",  
    "त्त्",   "त्त",     "क्त",  "दृ",  "कृ",
    "ह्न",  "ह्य",  "हृ",  "ह्म",  "ह्र",  "ह्",   "द्द",  "क्ष्", "क्ष", "त्र्", "त्र","ज्ञ",
    "छ्य",  "ट्य",  "ठ्य",  "ड्य",  "ढ्य", "द्य","द्व",
    "श्र",  "ट्र",    "ड्र",    "ढ्र",    "छ्र",   "क्र",  "फ्र",  "द्र",   "प्र",   "ग्र", "रु",  "रू",
    "्र",
    "ओ",  "औ",  "आ",   "अ",   "ई",   "इ",  "उ",   "ऊ",  "ऐ",  "ए", "ऋ",
    "क्",  "क",  "क्क",  "ख्",   "ख",    "ग्",   "ग",  "घ्",  "घ",    "ङ",
    "चै",   "च्",   "च",   "छ",  "ज्", "ज",   "झ्",  "झ",   "ञ",
    "ट्ट",   "ट्ठ",   "ट",   "ठ",   "ड्ड",   "ड्ढ",  "ड",   "ढ",  "ण्", "ण",  
    "त्",  "त",  "थ्", "थ",  "द्ध",  "द", "ध्", "ध",  "न्",  "न",  
    "प्",  "प",  "फ्", "फ",  "ब्",  "ब", "भ्",  "भ",  "म्",  "म",
    "य्",  "य",  "र",  "ल्", "ल",  "ळ",  "व्",  "व", 
    "श्", "श",  "ष्", "ष",  "स्",   "स",   "ह",     
    "ऑ",   "ॉ",  "ो",   "ौ",   "ा",   "ी",   "ु",   "ू",   "ृ",   "े",   "ै",
    "ं",   "ँ",   "ः",   "ॅ",    "ऽ",  "् ", "्" ]
    
    array_two = ["^", "*",  "Þ", "ß", "¼", "½", "¿", "À", "¾", "A", "\\", "&", "&", "Œ", "]","-","~ ", 
    "å",  "ƒ",  "„",   "…",   "†",   "‡",   "ˆ",   "‰",   "Š",   "‹","Û",
    "¶",   "d",    "[k",  "x",  "T",  "t",   "M+", "<+", "Q",  ";",    "j",   "u",
    "Ù",   "Ùk",   "Dr",    "–",   "—",       
    "à",   "á",    "â",   "ã",   "ºz",  "º",   "í", "{", "{k",  "«", "=","K", 
    "Nî",   "Vî",    "Bî",   "Mî",   "<î", "|","}",
    "J",   "Vª",   "Mª",  "<ªª",  "Nª",   "Ø",  "Ý",   "æ", "ç", "xz", "#", ":",
    "z",
    "vks",  "vkS",  "vk",    "v",   "bZ",  "b",  "m",  "Å",  ",s",  ",",   "_",
    "D",  "d",    "ô",     "[",     "[k",    "X",   "x",  "?",    "?k",   "³", 
    "pkS",  "P",    "p",  "N",   "T",    "t",   "÷",  ">",   "¥",
    "ê",      "ë",      "V",  "B",   "ì",       "ï",     "M",  "<",  ".", ".k",   
    "R",  "r",   "F", "Fk",  ")",    "n", "/",  "/k",  "U", "u",   
    "I",  "i",   "¶", "Q",   "C",  "c",  "H",  "Hk", "E",   "e",
    "¸",   ";",    "j",  "Y",   "y",  "G",  "O",  "o",
    "'", "'k",  "\"", "\"k", "L",   "l",   "g",      
    "v‚",    "‚",    "ks",   "kS",   "k",     "h",    "q",   "w",   "`",    "s",    "S",
    "a",    "¡",    "%",     "W",   "·",   "~ ", "~"]
    
    array_one_length = len(array_one)
    modified_substring = modified_substring.replace ("क़", "क़")   
    modified_substring = modified_substring.replace ("ख़‌", "ख़")
    modified_substring = modified_substring.replace ("ग़", "ग़")
    modified_substring = modified_substring.replace ("ज़", "ज़")
    modified_substring = modified_substring.replace ("ड़", "ड़")
    modified_substring = modified_substring.replace ("ढ़", "ढ़")
    modified_substring = modified_substring.replace ("ऩ", "ऩ")
    modified_substring = modified_substring.replace ("फ़", "फ़")
    modified_substring = modified_substring.replace ("य़", "य़")
    modified_substring = modified_substring.replace ("ऱ", "ऱ")
    modified_substring = modified_substring.replace("ि","f")
    
    for input_symbol_idx in range(0, array_one_length):
        modified_substring = modified_substring.replace(array_one[input_symbol_idx ] , array_two[input_symbol_idx] )
    
    modified_substring = "  " + modified_substring + "  "
    position_of_f = modified_substring.find("f")
    while (position_of_f != -1):    
        modified_substring = modified_substring[:position_of_f-1] + modified_substring[position_of_f] + modified_substring[position_of_f-1] + modified_substring[position_of_f+1:]
        position_of_f = modified_substring.find("f", position_of_f +1 )
    modified_substring = modified_substring.strip()
    
    modified_substring = "  " + modified_substring + "  "
    position_of_r = modified_substring.find("j~")
    set_of_matras =  ["‚",    "ks",   "kS",   "k",     "h",    "q",   "w",   "`",    "s",    "S", "a",    "¡",    "%",     "W",   "·",   "~ ", "~"]
    while (position_of_r != -1):    
        modified_substring = modified_substring.replace("j~","",1)
        if modified_substring[position_of_r + 1] in set_of_matras:
            modified_substring = modified_substring[:position_of_r + 2] + "Z" + modified_substring[position_of_r + 2:]
        else:
            modified_substring = modified_substring[:position_of_r + 1] + "Z" + modified_substring[position_of_r + 1:]
        position_of_r = modified_substring.find("j~")
    modified_substring = modified_substring.strip()
    
    return modified_substring

def extract_salutation_and_name(full_name, default_salutation="Mr."):
    if not full_name:
        return default_salutation, ""
    s = str(full_name).strip()
    s = re.sub(r'\s+', ' ', s)
    
    m = re.match(r'^((?:Mr\.?|Mrs\.?|Ms\.?|Shri|Smt\.?|Sh\.?))\b\s*(.*)', s, re.IGNORECASE)
    if m:
        sal = m.group(1).strip()
        sal_lower = sal.lower().rstrip('.')
        if sal_lower == "mr": sal = "Mr."
        elif sal_lower == "mrs": sal = "Mrs."
        elif sal_lower == "ms": sal = "Ms."
        elif sal_lower == "shri": sal = "Shri"
        elif sal_lower == "smt": sal = "Smt."
        elif sal_lower == "sh": sal = "Sh."
        return sal, m.group(2).strip()
    return default_salutation, s

def normalize_relative_salutation(name, relation):
    if not name:
        return ""
    s = str(name).strip()
    s = re.sub(r'\s+', ' ', s)
    
    if re.match(r'^(?:Mr\.?|Mrs\.?|Ms\.?|Shri|Smt\.?|Sh\.?)\b', s, re.IGNORECASE):
        m = re.match(r'^([a-z\.\:]+)\b(.*)', s, re.IGNORECASE)
        if m:
            sal = m.group(1).lower().rstrip('.')
            rest = m.group(2)
            if sal == "mr": s = "Mr." + rest
            elif sal == "mrs": s = "Mrs." + rest
            elif sal == "ms": s = "Ms." + rest
            elif sal == "shri": s = "Shri" + rest
            elif sal == "smt": s = "Smt." + rest
            elif sal == "sh": s = "Sh." + rest
        return s
        
    rel_lower = str(relation or "").strip().casefold()
    if any(x in rel_lower for x in ["mother", "m/o", "mother of", "husband", "husband of", "h/o"]) and not any(x in rel_lower for x in ["wife", "w/o", "father"]):
        s = "Mrs. " + s
    else:
        s = "Mr. " + s
    return s

def normalize_name_salutation(name, relation=None, default_to_male=True):
    if not name:
        return ""
    s = str(name).strip()
    s = re.sub(r'\s+', ' ', s)
    
    if re.match(r'^(?:Mr\.?|Mrs\.?|Ms\.?|Shri|Smt\.?|Sh\.?)\b', s, re.IGNORECASE):
        m = re.match(r'^([a-z\.\:]+)\b(.*)', s, re.IGNORECASE)
        if m:
            sal = m.group(1).lower().rstrip('.')
            rest = m.group(2)
            if sal == "mr": s = "Mr." + rest
            elif sal == "mrs": s = "Mrs." + rest
            elif sal == "ms": s = "Ms." + rest
            elif sal == "shri": s = "Shri" + rest
            elif sal == "smt": s = "Smt." + rest
            elif sal == "sh": s = "Sh." + rest
        return s
        
    is_female = False
    if relation:
        rel_lower = str(relation).strip().casefold()
        if any(x in rel_lower for x in ["w/o", "wife of", "wife", "mother", "m/o", "mother of", "d/o", "daughter of", "daughter"]):
            is_female = True
            
    if is_female:
        s = "Mrs. " + s
    else:
        if default_to_male:
            s = "Mr. " + s
    return s

LIST_DEFAULTS = {
    "bs": {"s": "", "n": "", "a": "", "r": "", "rn": "", "adr": "", "id": "", "pan": ""},
    "ls": {"n": "", "a": "", "w": "", "t": ""},
    "ps": {"adr": "", "n": "", "s": "", "e": "", "w": ""},
    "ws": {"n": "", "r": "", "rn": "", "adr": ""},
    "ds": {"t": ""},
}

HL_MARKER = "~~HL~~"

class TemplateProcessor:
    def __init__(self, template_path):
        self.template_path = template_path
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template not found at {template_path}")
        self.doc = DocxTemplate(template_path)

    def _convert_context_to_legacy(self, data):
        """Recursively scans context and encodes Hindi Unicode fields to DevLys ASCII."""
        if isinstance(data, dict):
            return {k: self._convert_context_to_legacy(v) for k, v in data.items()}
        elif isinstance(data, list):
            return [self._convert_context_to_legacy(x) for x in data]
        elif isinstance(data, str):
            # Convert only if contains Devanagari Hindi characters (U+0900 to U+097F)
            if any(0x0900 <= ord(c) <= 0x097F for c in data):
                return Unicode_to_KrutiDev(data)
            return data
        else:
            return data

    def _template_xml(self):
        parts = []
        with zipfile.ZipFile(self.template_path) as zf:
            for name in zf.namelist():
                if name.startswith("word/") and name.endswith(".xml"):
                    xml = zf.read(name).decode("utf-8", errors="ignore")
                    parts.append(re.sub(r"<[^>]+>", "", xml))
        return unescape("\n".join(parts))

    def _required_list_lengths(self):
        required = {}
        for list_name, index in re.findall(r"\b(bs|ls|ps|ws|ds)\s*\[\s*(\d+)\s*\]", self._template_xml()):
            required[list_name] = max(required.get(list_name, 0), int(index) + 1)
        return required

    def _pad_indexed_lists(self, context):
        for list_name, size in self._required_list_lengths().items():
            values = context.get(list_name)
            if not isinstance(values, list):
                values = []
            while len(values) < size:
                values.append(LIST_DEFAULTS[list_name].copy())
            context[list_name] = values
        return context

    def _normalize_context_salutations(self, context):
        # We need to normalize both the root context and the 'd' alias context if present
        contexts_to_clean = [context]
        if 'd' in context and isinstance(context['d'], dict):
            contexts_to_clean.append(context['d'])

        # Check if bank signatory salutation (bsign.s) is used anywhere in the template XML
        xml_content = self._template_xml()
        has_bsign_s = "bsign.s" in xml_content

        for ctx in contexts_to_clean:
            # 1. Normalize Borrowers relative names
            if "bs" in ctx and isinstance(ctx["bs"], list):
                for b in ctx["bs"]:
                    if isinstance(b, dict):
                        if b.get("rn"):
                            b["rn"] = normalize_relative_salutation(b["rn"], b.get("r"))
            
            # 2. Normalize Witnesses names and relative names
            if "ws" in ctx and isinstance(ctx["ws"], list):
                for w in ctx["ws"]:
                    if isinstance(w, dict):
                        if w.get("n"):
                            w["n"] = normalize_name_salutation(w["n"], w.get("r"))
                        if w.get("rn"):
                            w["rn"] = normalize_relative_salutation(w["rn"], w.get("r"))

            # 3. Normalize Bank Signatory
            if "bsign" in ctx and isinstance(ctx["bsign"], dict):
                bsign = ctx["bsign"]
                # Normalize signatory relative name
                if bsign.get("rn"):
                    bsign["rn"] = normalize_relative_salutation(bsign["rn"], bsign.get("r"))
                
                # Normalize name and split or embed salutation depending on has_bsign_s
                if bsign.get("n"):
                    raw_name = bsign["n"]
                    sal, clean_name = extract_salutation_and_name(raw_name)
                    
                    if has_bsign_s:
                        bsign["s"] = sal
                        bsign["n"] = clean_name
                    else:
                        bsign["s"] = ""
                        bsign["n"] = normalize_name_salutation(raw_name, bsign.get("r"))

    def generate(self, context, output_path, highlight_ai=False, highlight_missing=False, verified_fields=None):
        """
        context: A dictionary containing the data to fill in the template.
        output_path: Where to save the generated .docx file.
        verified_fields: A set of field paths (e.g., "bs.0.n") that are verified.
        """
        if verified_fields is None: verified_fields = set()
        
        # Globally normalize all salutations dynamically in the context
        self._normalize_context_salutations(context)
        
        # Ensure 'd' alias is padded correctly if it exists, otherwise pad root context
        data_to_pad = context.get('d', context)
        self._pad_indexed_lists(data_to_pad)

        # Apply markers for highlighting if requested
        if highlight_ai or highlight_missing:
            context = self._apply_highlight_markers(context, verified_fields, highlight_ai, highlight_missing)

        # Convert any Unicode Hindi values to legacy DevLys ASCII characters before rendering
        context = self._convert_context_to_legacy(context)

        # Render placeholders with markers (markers are treated as plain text)
        self.doc.render(context)

        # Post-process markers in-memory before saving.
        # We ONLY process the body and tables to avoid header/footer corruption.
        if highlight_ai or highlight_missing:
            self._apply_body_highlights()

        self.doc.save(output_path)
        return output_path

    def _apply_highlight_markers(self, data, verified, h_ai, h_miss, path=""):
        if isinstance(data, dict):
            new_data = {}
            for k, v in data.items():
                new_path = f"{path}.{k}" if path else k
                new_data[k] = self._apply_highlight_markers(v, verified, h_ai, h_miss, new_path)
            return new_data
        elif isinstance(data, list):
            return [self._apply_highlight_markers(item, verified, h_ai, h_miss, f"{path}.{i}") for i, item in enumerate(data)]
        else:
            val = str(data).strip()
            if not val:
                if h_miss:
                    return f"{HL_MARKER}[MISSING]{HL_MARKER}"
                return ""

            # If not verified and h_ai is True, wrap in markers
            if h_ai and path not in verified:
                return f"{HL_MARKER}{val}{HL_MARKER}"
            
            return val

    def _apply_body_highlights(self):
        """Surgically applies highlights to the main body only, skipping headers/footers."""
        
        def highlight_in_container(container):
            # 1. Process all paragraphs in this container (including table cells)
            for p in container.paragraphs:
                for run in p.runs:
                    if HL_MARKER in run.text:
                        run.text = run.text.replace(HL_MARKER, "")
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            
            # 2. Recursively check tables within this container
            if hasattr(container, 'tables'):
                for table in container.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            highlight_in_container(cell)

        # Run exclusively on the main document body
        highlight_in_container(self.doc)

if __name__ == "__main__":
    # Small test if a test_template.docx exists
    pass
