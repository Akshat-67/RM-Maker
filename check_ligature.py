import docx
doc = docx.Document('Validated_AI.docx')
p2 = doc.paragraphs[2].text
# find 'mRrjkf/kdkjh' or similar
import re
m = re.search(r'm[RrÙ]+jkf/kdkjh', p2)
if m:
    s = m.group(0)
    print(f"GEN SYM: '{s}'")
    print(f"HEX: {' '.join(hex(ord(c)) for c in s)}")

doc_act = docx.Document('Test/ACTUAL.docx')
p2_act = [p.text for p in doc_act.paragraphs if 'vkt fnukad' in p.text][0]
m_act = re.search(r'm[RrÙ]+jkf/kdkjh', p2_act)
if m_act:
    s_act = m_act.group(0)
    print(f"ACT SYM: '{s_act}'")
    print(f"HEX: {' '.join(hex(ord(c)) for c in s_act)}")
