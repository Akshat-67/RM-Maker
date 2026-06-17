import docx

def fix_template_drafter():
    path = "templates/SALE_DEED/SD-Vivek Saxena,  Sunita Saxena - Vijay Laxmi - JDA+2SD+Flat_unicode (1)_devlys.docx"
    doc = docx.Document(path)
    
    # 1. Update the 'para' placeholder to use Jinja2 multi-paragraph loop
    # This is safer than doing it in code for every document.
    for p in doc.paragraphs:
        if "{{para}}" in p.text:
            p.text = p.text.replace("{{para}}", "{% p for line in paragraphs %}{{line}}{% p endfor %}")
            print("Template updated with multi-paragraph Jinja2 loop.")
            
    # 2. Remove redundant dimension paragraph (P6 in original list)
    # We identify it by its specific placeholders
    to_delete = []
    for p in doc.paragraphs:
        if "{{ps[0].dimension_text}}" in p.text and "{{ps[0].boundary_text}}" in p.text:
            to_delete.append(p)
            print("Redundant dimension paragraph marked for deletion.")
            
    for p in to_delete:
        p_element = p._p
        p_element.getparent().remove(p_element)
        
    doc.save(path)

if __name__ == "__main__":
    fix_template_drafter()
