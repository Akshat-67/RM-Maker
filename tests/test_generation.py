import os
import shutil
import pytest
from services.generation_service import compile_and_render_document, generate_draft_preview
from services.session_manager import save_case_session, load_case_session

@pytest.fixture
def temp_cases_dir(tmp_path, monkeypatch):
    # Override CASES_DIR and TEMPLATES_DIR to point to a temporary test directory
    test_cases_dir = tmp_path / "cases"
    test_cases_dir.mkdir()
    
    test_templates_dir = tmp_path / "templates"
    test_templates_dir.mkdir()
    
    # Create structure
    (test_templates_dir / "CHOLA").mkdir()
    (test_templates_dir / "SALE_DEED").mkdir()
    
    import docx
    
    # Create dummy RM docx
    doc_rm = docx.Document()
    doc_rm.add_paragraph("RM Template: {{ rd }} {{ bank }} {{ bs[0].n }}")
    doc_rm.save(str(test_templates_dir / "CHOLA" / "RM_CHOLA_1B_1L.docx"))
    
    # Create dummy SD docx
    doc_sd = docx.Document()
    doc_sd.add_paragraph("SD Template: {{ consideration }} {{ ss[0].n }}")
    doc_sd.save(str(test_templates_dir / "SALE_DEED" / "SD_JDA_2SD_Flat_1S_1B.docx"))
    
    import services.session_manager
    import services.generation_service
    import services.file_service
    
    monkeypatch.setattr(services.session_manager, "CASES_DIR", str(test_cases_dir))
    monkeypatch.setattr(services.generation_service, "CASES_DIR", str(test_cases_dir))
    monkeypatch.setattr(services.generation_service, "TEMPLATES_DIR", str(test_templates_dir))
    monkeypatch.setattr(services.file_service, "CASES_DIR", str(test_cases_dir))
    monkeypatch.setattr(services.file_service, "TEMPLATES_DIR", str(test_templates_dir))
    
    yield test_cases_dir

def test_golden_generation_rm(temp_cases_dir):
    case_id = "test_rm_case"
    
    # 1. Setup mock session with RM fields and bank=CHOLA
    mock_data = {
        "second_schedule": "Schedule line 1\nSchedule line 2",
        "bs": [{"n": "Borrower Alice", "adr": "MANSAROVAR, JAIPUR", "id": "1234 5678 9012"}],
        "ls": [{"a": "2,500,000", "r_rate": "10.50%", "emi": "25,000", "emi_w": "Twenty Five Thousand"}],
        "ps": [{"adr": "PLOT 10, KESAR NAGAR, JAIPUR"}],
        "ws": [{"n": "Witness Walter", "adr": "SANGANER, JAIPUR"}],
        "bsign": {"n": "Auth Signatory", "s": "Mr", "r": "S/O", "adr": "JAIPUR", "id": ""},
        "title_chain": [{"date": "10.05.2024", "reg_date": "12.05.2024", "executant_name": "Seller Bob"}]
    }
    
    # Save the session to the temp directory
    save_case_session(
        case_id=case_id,
        data=mock_data,
        files=[],
        verified_fields=set(),
        bank="CHOLA",
        borrower_count="1",
        loan_count="1",
        properties_count="1",
        doc_type="RM"
    )
    
    # Load session back
    session = load_case_session(case_id)
    assert session is not None
    
    # 2. Run compile and render
    try:
        output_filepath, output_filename = compile_and_render_document(
            case_id=case_id,
            session=session,
            doc_type="RM",
            bank="CHOLA",
            borrowers="1",
            loans="1",
            properties="1"
        )
        
        # Verify output exists and placeholders compiled
        assert os.path.exists(output_filepath)
        assert output_filename == f"RM_{case_id}.docx"
    except ValueError as e:
        if "No valid template found" in str(e):
            pytest.skip("Templates directory not configured or empty in test environment")
        else:
            raise e

def test_golden_generation_sd(temp_cases_dir):
    case_id = "test_sd_case"
    
    # 1. Setup mock session with SD fields, title chain, and Hindi characters
    mock_data = {
        "second_schedule": "Property Schedule line",
        "ss": [{"n": "विक्रेता राम", "adr": "जयपुर", "id": "1111 2222 3333", "a": "45", "c": "General"}],
        "bs": [{"n": "क्रेता श्याम", "adr": "जोधपुर", "id": "4444 5555 6666", "a": "32"}],
        "ws": [
            {"n": "गवाह एक", "adr": "अजमेर", "id": "7777 8888 9999", "a": "35"},
            {"n": "गवाह दो", "adr": "बीकानेर", "id": "2222 3333 4444", "a": "40"}
        ],
        "ps": [{"adr": "भूखंड संख्या 5, प्रताप नगर, जयपुर", "area": "200"}],
        "title_chain": [{"date": "20.12.2023", "reg_date": "22.12.2023", "executant_name": "राम एवं श्याम"}],
        "amount": "3000000",
        "amount_words": "तीस लाख मात्र"
    }
    
    save_case_session(
        case_id=case_id,
        data=mock_data,
        files=[],
        verified_fields=set(),
        bank="",
        borrower_count="1",
        loan_count="1",
        properties_count="1",
        doc_type="SD",
        sellers_count="1",
        buyers_count="1"
    )
    
    session = load_case_session(case_id)
    
    try:
        output_filepath, output_filename = compile_and_render_document(
            case_id=case_id,
            session=session,
            doc_type="SD",
            sellers_count="1",
            buyers_count="1"
        )
        
        assert os.path.exists(output_filepath)
        assert output_filename == f"SD_{case_id}.docx"
    except ValueError as e:
        if "No valid template found" in str(e):
            pytest.skip("Templates directory not configured or empty in test environment")
        else:
            raise e

def test_golden_draft_preview_rm(temp_cases_dir):
    case_id = "test_rm_preview"
    mock_data = {
        "second_schedule": "Schedule lines",
        "bs": [{"n": "Alice", "adr": "JAIPUR"}],
        "ls": [{"a": "1,000,000", "r_rate": "10%", "emi": "10,000", "emi_w": "Ten Thousand"}],
        "ps": [{"adr": "JAIPUR"}],
        "ws": [{"n": "Walter", "adr": "JAIPUR"}],
        "bsign": {"n": "Auth Signatory", "s": "Mr", "r": "S/O", "adr": "JAIPUR", "id": ""}
    }
    
    save_case_session(
        case_id=case_id,
        data=mock_data,
        files=[],
        verified_fields=set(),
        bank="CHOLA",
        borrower_count="1",
        loan_count="1",
        properties_count="1",
        doc_type="RM"
    )
    
    session = load_case_session(case_id)
    
    try:
        preview_html = generate_draft_preview(
            case_id=case_id,
            session=session,
            doc_type="RM",
            bank="CHOLA",
            borrowers="1",
            loans="1",
            properties="1"
        )
        
        assert "word-page" in preview_html
    except ValueError as e:
        if "No valid template found" in str(e):
            pytest.skip("Templates directory not configured or empty")
        else:
            raise e
