import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document("templates/SALE_DEED/SD_JDA_2SD_Flat_2S_1B.docx")

print("--- SEARCHING PARAGRAPHS FOR SELLERS / BUYERS / WITNESSES ---")
keywords = ["foosd", "lquhrk", "fot;", "lqYrku", "iou"]

for idx, p in enumerate(doc.paragraphs):
    text = p.text
    if any(k in text for k in keywords):
        print(f"\n[Para {idx}]: {text}")

for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            text = cell.text
            if any(k in text for k in keywords):
                print(f"\n[Table {t_idx}, Row {r_idx}, Cell {c_idx}]: {text}")
