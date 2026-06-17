import json
import os
import docx
from modules.sd.processor import SDTemplateProcessor
from modules.sd.extractor import SDDataExtractor
from modules.sd.narrative import generate_chain_narrative

def debug_generate():
    case_id = "case_1781374115"
    case_path = os.path.join("cases", case_id, "session.json")
    with open(case_path, "r", encoding="utf-8") as f:
        session = json.load(f)
    
    data = session["data"]
    doc_type = "SD"
    property_type = session.get("property_type", "Plot")
    
    from collections import defaultdict
    for key in ["bs", "ls", "ps", "ws", "ss", "sellers", "buyers", "chain"]:
        if key not in data or not isinstance(data[key], list):
            data[key] = []
        while len(data[key]) < 10:
            data[key].append(dict())

    context = data.copy()
    context["w1"] = data["ws"][0]
    context["w2"] = data["ws"][1]
    
    # SD-specific field aliasing
    for list_key in ["ss", "bs", "ws"]:
        for person in context.get(list_key, []):
            if isinstance(person, dict):
                if "address" not in person or not person["address"]:
                    person["address"] = person.get("adr", "")
                if "aadhaar" not in person or not person["aadhaar"]:
                    person["aadhaar"] = person.get("id", "")
                    
    # Recompute ps
    _sd_extractor = SDDataExtractor()
    for p in context.get("ps", []):
        if isinstance(p, dict) and p:
            p["full_address"] = _sd_extractor.generate_full_property_address(p, "SD", property_type)
            p["dimension_text"] = _sd_extractor.generate_dimension_text(p)
            p["boundary_text"] = _sd_extractor.generate_boundary_text(p)
            
    if "title_chain" in context and isinstance(context["title_chain"], list):
        ps0 = context.get("ps", [{}])[0]
        chain_paras = generate_chain_narrative(context["title_chain"], property_details=ps0, context=context)
        context["chain_paragraphs"] = chain_paras
        context["chain_text"] = "\n\n\tतत्पश्चात् ".join(chain_paras)
        
    context['d'] = context.copy()
    
    template_path = "templates/SALE_DEED/SD-Vivek Saxena,  Sunita Saxena - Vijay Laxmi - JDA+2SD+Flat_unicode (1)_devlys.docx"
    output_path = "ReEval_AI.docx"
    
    processor = SDTemplateProcessor(template_path)
    # Enable highlighting
    processor.generate(context, output_path, highlight_ai=True, highlight_missing=True)
    print(f"Generated {output_path}")
    
    # Print out paragraphs for review
    doc_actual = docx.Document("Test/ACTUAL.docx")
    doc_ai = docx.Document("ReEval_AI.docx")
    
    print("\n\n--- ACTUAL.docx (First 15 Paragraphs) ---")
    for i, p in enumerate(doc_actual.paragraphs[:15]):
        text = p.text.strip()
        if text:
            print(f"[{i}]: {text}")
            
    print("\n\n--- ReEval_AI.docx (First 15 Paragraphs) ---")
    for i, p in enumerate(doc_ai.paragraphs[:15]):
        text = p.text.strip()
        if text:
            print(f"[{i}]: {text}")
            
if __name__ == "__main__":
    debug_generate()
