import docx
import difflib

def compare_docs(actual_path, ai_path, output_diff_path="parity_diff_current.txt"):
    doc_actual = docx.Document(actual_path)
    doc_ai = docx.Document(ai_path)
    
    # Filter out empty paragraphs to focus on content
    actual_paras = [p.text.strip() for p in doc_actual.paragraphs if p.text.strip()]
    ai_paras = [p.text.strip() for p in doc_ai.paragraphs if p.text.strip()]
    
    out = []
    # Comparison loop
    max_len = max(len(actual_paras), len(ai_paras))
    for i in range(max_len):
        act = actual_paras[i] if i < len(actual_paras) else "[MISSING PARA]"
        gen = ai_paras[i] if i < len(ai_paras) else "[MISSING PARA]"
        
        if act != gen:
            out.append(f"\n--- PARA {i} MISMATCH ---")
            out.append(f"ACTUAL: {act}")
            out.append(f"GEN   : {gen}")
            
            # Show character level diff for clarity
            d = difflib.Differ()
            diff = list(d.compare([act], [gen]))
            # This is simpler for character level:
            res = ""
            for j, s in enumerate(difflib.ndiff(act, gen)):
                if s[0] == ' ': continue
                elif s[0] == '-': res += f" (ACTUAL_ONLY: '{s[-1]}')"
                elif s[0] == '+': res += f" (GEN_ONLY: '{s[-1]}')"
            if res:
                out.append(f"DIFF  : {res}")
                
    with open(output_diff_path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))
    print(f"Generated {output_diff_path} with {len(out)} lines of diff.")

if __name__ == "__main__":
    compare_docs("Test/ACTUAL.docx", "Validated_AI.docx")
