import os
import re
import json
from docx import Document
from services.ocr_service import ocr_extract_text
from services.ai_client import AIClient
from services.session_manager import load_case_session

CASES_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "cases")

def run_case_audit(case_id, doc_type="RM"):
    """
    Runs rule-based heuristics and Gemini-powered semantic auditing
    on the generated draft text against the raw OCR cache.
    """
    session = load_case_session(case_id)
    if not session:
        return {"error": f"Case session '{case_id}' not found"}
        
    discrepancies = []
    
    # 1. Local Rule-based audits (100% Free and Offline)
    draft_dir = os.path.join(CASES_DIR, case_id)
    if not os.path.exists(draft_dir):
        return {"discrepancies": []}
        
    draft_files = [os.path.join(draft_dir, f) for f in os.listdir(draft_dir) if f.lower().endswith(".docx") and not f.startswith("~")]
    
    draft_text = ""
    if draft_files:
        try:
            doc = Document(draft_files[0])
            draft_text = "\n".join([p.text for p in doc.paragraphs])
            # Append tables text to draft text
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    draft_text += "\n" + row_text
        except Exception:
            pass

    if not draft_text:
        return {"discrepancies": []}

    # Check for placeholder leaks
    if "{{" in draft_text or "}}" in draft_text:
        discrepancies.append({
            "id": "leak_placeholders",
            "type": "layout_error",
            "message": "Unreplaced Jinja placeholders ({{ or }}) detected in generated draft document.",
            "field_path": "",
            "suggested_fix": "",
            "autofixable": False
        })

    # Check for duplicate words
    dup_match = re.search(r'\b(the|and|or|of|to|is|a)\s+\1\b', draft_text, re.IGNORECASE)
    if dup_match:
        discrepancies.append({
            "id": "duplicate_words",
            "type": "format_error",
            "message": f"Duplicated words '{dup_match.group(0)}' found in draft text.",
            "field_path": "",
            "suggested_fix": "",
            "autofixable": False
        })

    # Check missing relative/salutations in RM context (English names only)
    data = session.get("data", {})
    for idx, b in enumerate(data.get("bs", [])):
        name = str(b.get("n", "") or "").strip()
        salutation = str(b.get("s", "") or "").strip()
        
        # A salutation is already present if the separate 's' field is set, or if the name starts with one
        has_salutation = bool(salutation) or any(name.startswith(x) for x in ["Mr. ", "Mrs. ", "Ms. "])
        
        if name and not has_salutation:
            # Determine appropriate salutation based on relation
            relation = str(b.get("r", "") or "").strip().lower()
            if relation in ["w/o", "wife of"]:
                sugg_prefix = "Mrs."
            elif relation in ["d/o", "daughter of"]:
                sugg_prefix = "Ms."
            else:
                sugg_prefix = "Mr."
                
            discrepancies.append({
                "id": f"missing_salutation_bs_{idx}",
                "type": "missing_salutation",
                "message": f"Borrower {idx + 1} '{name}' is missing a salutation (Mr. / Mrs. / Ms.).",
                "field_path": f"bs.{idx}.n",
                "suggested_fix": f"{sugg_prefix} {name}",
                "autofixable": True
            })

    # Check missing salutation for Bank Signatory (RM context)
    bsign = data.get("bsign", {})
    if bsign:
        name = str(bsign.get("n", "") or "").strip()
        salutation = str(bsign.get("s", "") or "").strip()
        has_salutation = bool(salutation) or any(name.startswith(x) for x in ["Mr. ", "Mrs. ", "Ms. "])
        
        if name and not has_salutation:
            relation = str(bsign.get("r", "") or "").strip().lower()
            if relation in ["w/o", "wife of"]:
                sugg_prefix = "Mrs."
            elif relation in ["d/o", "daughter of"]:
                sugg_prefix = "Ms."
            else:
                sugg_prefix = "Mr."
                
            discrepancies.append({
                "id": "missing_salutation_bsign",
                "type": "missing_salutation",
                "message": f"Bank Signatory '{name}' is missing a salutation (Mr. / Mrs. / Ms.).",
                "field_path": "bsign.n",
                "suggested_fix": f"{sugg_prefix} {name}",
                "autofixable": True
            })

    # 2. Semantic Gemini Audit (English Only)
    # Fetch OCR texts from all uploaded KYC files
    uploaded_files = [os.path.join(draft_dir, f) for f in session.get("files", [])]
    ocr_text = ocr_extract_text(case_id, uploaded_files)
    
    if ocr_text and draft_text:
        prompt = f"""
        You are a senior legal compliance auditor. Compare the source document text against the compiled draft document.
        Spot any inaccuracies, mismatched names, mismatched borrower settings, or key details missing.
        
        Since this is a Registered Mortgage (RM) case, all text and suggested fixes MUST be strictly in English.
        Do NOT suggest any Hindi text.
        
        NOTE: In our data model, borrower/signatory salutations (Mr./Mrs./Ms.) are stored in a separate dropdown field 's', and the clean name is stored in 'n'. If suggesting a fix for a borrower name, target the name field 'bs.0.n' and provide the clean name WITHOUT the prefix, or provide it with the prefix (e.g. 'Mr. John Doe') and our backend will automatically parse and split it.
        
        SOURCE DOCUMENTS GROUND TRUTH:
        {ocr_text}
        
        COMPILED DRAFT TEXT:
        {draft_text}
        
        Respond ONLY with a JSON object containing a "discrepancies" array of issues:
        {{
            "discrepancies": [
                {{
                    "id": "mismatch_borrower_name",
                    "type": "mismatch",
                    "message": "Detail mismatch description in English",
                    "field_path": "bs.0.n",
                    "suggested_fix": "Clean English suggested value",
                    "autofixable": true
                }}
            ]
        }}
        """
        try:
            client = AIClient()
            raw_res = client.generate_text(prompt, model="gemini-2.5-flash")
            match = re.search(r'\{.*\}', raw_res, re.DOTALL)
            if match:
                res_data = json.loads(match.group(0))
                discrepancies.extend(res_data.get("discrepancies", []))
        except Exception:
            pass

    return {"discrepancies": discrepancies}
